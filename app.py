import base64
import calendar
import hashlib
import hmac
import os
import re
import secrets
import time
import unicodedata
from datetime import date, datetime
from html import escape
from io import BytesIO
from zoneinfo import ZoneInfo

import pandas as pd
import psycopg2
import streamlit as st
from database import (
    COLUMNAS_REGISTRO,
    ITERACIONES_HASH_CONTRASENA,
    USUARIOS_PREDEFINIDOS,
    actualizar_registro,
    actualizar_ultimo_inicio_sesion,
    conectar,
    consultar_registros,
    consultar_usuarios,
    consultar_usuarios_eliminados,
    eliminar_registro,
    eliminar_registros_supervisor_en_conexion,
    guardar_usuarios as guardar_usuarios_en_db,
    guardar_usuarios_en_conexion,
    inicializar_base_datos,
    insertar_registro,
    normalizar_usuarios,
)


# Configuración de la página
st.set_page_config(
    page_title="El Búnker - Sistema de Autolavado",
    layout="wide",
    initial_sidebar_state="expanded",
)
with open(os.path.join(os.path.dirname(__file__), "fondo_sidebar.png"), "rb") as archivo_fondo:
    fondo_sidebar_base64 = base64.b64encode(archivo_fondo.read()).decode("ascii")
with open(os.path.join(os.path.dirname(__file__), "bunker_logo.png"), "rb") as archivo_logo:
    logo_base64 = base64.b64encode(archivo_logo.read()).decode("ascii")
with open(os.path.join(os.path.dirname(__file__), "banner_bunker.png"), "rb") as archivo_banner:
    banner_base64 = base64.b64encode(archivo_banner.read()).decode("ascii")

ROLES_USUARIO = ["Supervisor", "Administrador"]
FECHA_MINIMA_CALENDARIO = date(2026, 1, 1)
FECHA_MAXIMA_CALENDARIO = date(2036, 12, 31)
ZONA_HORARIA_PERU = ZoneInfo("America/Lima")


def construir_titulo_seccion(seccion):
    texto = unicodedata.normalize("NFKD", f"{seccion} - BUNKER")
    return texto.encode("ascii", "ignore").decode("ascii").upper()


def obtener_fecha_selector(clave, fecha_predeterminada):
    fecha = st.session_state.get(clave)
    if isinstance(fecha, datetime):
        fecha = fecha.date()
    if (
        not isinstance(fecha, date)
        or not FECHA_MINIMA_CALENDARIO
        <= fecha
        <= FECHA_MAXIMA_CALENDARIO
    ):
        st.session_state.pop(clave, None)
        fecha = fecha_predeterminada
    return min(max(fecha, FECHA_MINIMA_CALENDARIO), FECHA_MAXIMA_CALENDARIO)


def actualizar_fecha_calendario(clave, fecha, clave_mes, clave_anio, on_change):
    fecha_anterior = obtener_fecha_selector(clave, fecha)
    st.session_state[clave] = fecha
    st.session_state[clave_mes] = fecha.month
    st.session_state[clave_anio] = fecha.year
    if fecha != fecha_anterior and on_change is not None:
        on_change()


def seleccionar_fecha(etiqueta, clave, fecha_predeterminada, on_change=None):
    meses = (
        "Enero",
        "Febrero",
        "Marzo",
        "Abril",
        "Mayo",
        "Junio",
        "Julio",
        "Agosto",
        "Septiembre",
        "Octubre",
        "Noviembre",
        "Diciembre",
    )
    fecha_seleccionada = obtener_fecha_selector(clave, fecha_predeterminada)
    clave_mes = f"{clave}_mes_calendario"
    clave_anio = f"{clave}_anio_calendario"
    clave_popover = f"{clave}_popover_calendario"
    st.session_state.setdefault(clave_mes, fecha_seleccionada.month)
    st.session_state.setdefault(clave_anio, fecha_seleccionada.year)

    st.markdown(
        f'<div style="margin-bottom:0.25rem;font-size:0.875rem;">'
        f"{escape(etiqueta)}</div>",
        unsafe_allow_html=True,
    )
    with st.popover(
        f":material/calendar_month: {fecha_seleccionada:%d/%m/%Y}",
        key=clave_popover,
        width=340,
    ):
        with st.container(width=300):
            mes_visible = st.session_state[clave_mes]
            anio_visible = st.session_state[clave_anio]
            columnas_semana = st.columns(7, gap="small")
            for columna, dia_semana in zip(
                columnas_semana, ("Lu", "Ma", "Mi", "Ju", "Vi", "Sa", "Do")
            ):
                with columna:
                    st.markdown(
                        f'<div style="text-align:center;color:#9eaaa0;'
                        f'font-size:0.75rem;font-weight:600;">{dia_semana}</div>',
                        unsafe_allow_html=True,
                    )
            semanas = calendar.Calendar(firstweekday=0).monthdayscalendar(
                anio_visible, mes_visible
            )
            for semana in semanas:
                columnas_dias = st.columns(7, gap="small")
                for columna, numero_dia in zip(columnas_dias, semana):
                    with columna:
                        if numero_dia == 0:
                            st.markdown("&nbsp;", unsafe_allow_html=True)
                            continue
                        fecha_dia = date(anio_visible, mes_visible, numero_dia)
                        st.button(
                            str(numero_dia),
                            key=f"{clave}_dia_{fecha_dia.isoformat()}",
                            type=(
                                "primary"
                                if fecha_dia == fecha_seleccionada
                                else "secondary"
                            ),
                            use_container_width=True,
                            on_click=actualizar_fecha_calendario,
                            args=(clave, fecha_dia, clave_mes, clave_anio, on_change),
                        )

            columnas_navegacion = st.columns(2, gap="small")
            with columnas_navegacion[0]:
                st.selectbox(
                    "Mes",
                    options=range(1, 13),
                    format_func=lambda numero: meses[numero - 1],
                    key=clave_mes,
                    label_visibility="collapsed",
                )
            with columnas_navegacion[1]:
                st.selectbox(
                    "Año",
                    options=range(
                        FECHA_MINIMA_CALENDARIO.year,
                        FECHA_MAXIMA_CALENDARIO.year + 1,
                    ),
                    key=clave_anio,
                    label_visibility="collapsed",
                )

    return fecha_seleccionada


def verificar_contrasena(cuenta, contrasena):
    if not cuenta["salt"] or not cuenta["contrasena_hash"]:
        return False
    hash_ingresado = hashlib.pbkdf2_hmac(
        "sha256",
        contrasena.encode("utf-8"),
        bytes.fromhex(cuenta["salt"]),
        ITERACIONES_HASH_CONTRASENA,
    ).hex()
    return hmac.compare_digest(cuenta["contrasena_hash"], hash_ingresado)


def cargar_usuarios():
    if (
        "_usuarios_sesion" not in st.session_state
        or "_usuarios_eliminados_sesion" not in st.session_state
    ):
        st.session_state["_usuarios_sesion"] = consultar_usuarios()
        st.session_state["_usuarios_eliminados_sesion"] = (
            consultar_usuarios_eliminados()
        )
    return (
        st.session_state["_usuarios_sesion"],
        st.session_state["_usuarios_eliminados_sesion"],
    )


def _actualizar_usuarios_en_sesion(
    usuarios_actualizados, eliminados_actualizados
):
    global usuarios, usuarios_eliminados_definitivamente

    usuarios = {
        nombre: cuenta.copy()
        for nombre, cuenta in usuarios_actualizados.items()
    }
    usuarios_eliminados_definitivamente = set(eliminados_actualizados)
    st.session_state["_usuarios_sesion"] = usuarios
    st.session_state["_usuarios_eliminados_sesion"] = (
        usuarios_eliminados_definitivamente
    )


def guardar_usuarios(
    usuarios_actualizados, usuarios_eliminados_definitivamente=()
):
    usuarios_normalizados = normalizar_usuarios(usuarios_actualizados)
    guardar_usuarios_en_db(
        usuarios_normalizados,
        usuarios_eliminados_definitivamente,
    )
    _actualizar_usuarios_en_sesion(
        usuarios_normalizados, usuarios_eliminados_definitivamente
    )


@st.dialog("Agregar nuevo usuario")
def mostrar_formulario_nuevo_usuario():
    st.markdown("Registra una cuenta con uno de los roles disponibles.")
    with st.form("form_nuevo_usuario", clear_on_submit=True):
        nuevo_nombre = st.text_input("Nombre de usuario")
        nueva_contrasena = st.text_input("Contraseña", type="password")
        confirmar_contrasena = st.text_input(
            "Confirmar contraseña", type="password"
        )
        nuevo_rol = st.selectbox("Rol", ROLES_USUARIO)
        guardar_nuevo_usuario = st.form_submit_button(
            "Guardar y agregar usuario",
            type="primary",
            use_container_width=True,
        )
    if not guardar_nuevo_usuario:
        return

    nombre_limpio = nuevo_nombre.strip()
    if not nombre_limpio:
        st.warning("Escribe un nombre de usuario.")
    elif any(
        nombre_limpio.casefold() == existente.casefold()
        for existente in usuarios
    ):
        st.warning("Ese nombre de usuario ya está registrado.")
    elif any(
        nombre_limpio.casefold() == predefinido.casefold()
        for predefinido in USUARIOS_PREDEFINIDOS
    ):
        st.warning("Ese nombre está reservado y no se puede volver a registrar.")
    elif len(nueva_contrasena) < 8:
        st.warning("La contraseña debe tener al menos 8 caracteres.")
    elif nueva_contrasena != confirmar_contrasena:
        st.warning("Las contraseñas no coinciden.")
    else:
        salt = secrets.token_bytes(16)
        cuenta_nueva = {
            "contrasena": None,
            "contrasena_hash": hashlib.pbkdf2_hmac(
                "sha256",
                nueva_contrasena.encode("utf-8"),
                salt,
                ITERACIONES_HASH_CONTRASENA,
            ).hex(),
            "salt": salt.hex(),
            "rol": nuevo_rol,
            "ultimo_inicio_sesion": None,
            "activo": True,
            "predefinido": False,
        }
        usuarios_actualizados = dict(usuarios)
        usuarios_actualizados[nombre_limpio] = cuenta_nueva
        try:
            guardar_usuarios(
                usuarios_actualizados,
                usuarios_eliminados_definitivamente=(
                    usuarios_eliminados_definitivamente
                ),
            )
        except (psycopg2.Error, ValueError) as error:
            st.error(f"No se pudo guardar el nuevo usuario: {error}")
        else:
            st.session_state["mensaje_gestion_usuarios"] = (
                f"El usuario {nombre_limpio} fue agregado correctamente."
            )
            st.rerun()


@st.dialog("Editar usuario")
def mostrar_formulario_editar_usuario():
    cuentas_activas = [
        nombre for nombre, cuenta in usuarios.items() if cuenta["activo"]
    ]
    if not cuentas_activas:
        st.info("No hay cuentas activas para editar.")
        return

    usuario_seleccionado = st.selectbox(
        "Seleccionar usuario",
        cuentas_activas,
        key="usuario_a_editar",
    )
    cuenta_seleccionada = usuarios[usuario_seleccionado]
    clave_usuario = hashlib.sha256(
        usuario_seleccionado.casefold().encode("utf-8")
    ).hexdigest()[:12]
    rol_actual = cuenta_seleccionada["rol"]
    rol_elegido = st.selectbox(
        "Rol",
        ROLES_USUARIO,
        index=ROLES_USUARIO.index(
            rol_actual if rol_actual in ROLES_USUARIO else "Supervisor"
        ),
        key=f"rol_editar_usuario_{clave_usuario}",
    )
    guardar_cambios = st.button(
        "Guardar cambios",
        type="primary",
        use_container_width=True,
    )
    if not guardar_cambios:
        return

    usuarios_actualizados = {
        nombre: cuenta.copy() for nombre, cuenta in usuarios.items()
    }
    usuarios_actualizados[usuario_seleccionado]["rol"] = rol_elegido
    try:
        guardar_usuarios(
            usuarios_actualizados,
            usuarios_eliminados_definitivamente=(
                usuarios_eliminados_definitivamente
            ),
        )
    except (psycopg2.Error, ValueError) as error:
        st.error(f"No se pudo guardar el usuario: {error}")
    else:
        st.session_state["mensaje_gestion_usuarios"] = (
            f"Los datos de {usuario_seleccionado} fueron actualizados."
        )
        st.rerun()


@st.dialog("Dar de baja una cuenta")
def mostrar_formulario_baja_usuario():
    st.markdown("Selecciona la cuenta que deseas dar de baja y confirma con tu contraseña de administrador.")
    cuentas_eliminables = [
        nombre
        for nombre, cuenta in usuarios.items()
        if cuenta["activo"]
        and cuenta["rol"] != "Administrador"
        and nombre != st.session_state.usuario_actual
    ]
    with st.form("form_baja_usuario", clear_on_submit=True):
        usuario_seleccionado = st.selectbox(
            "Seleccionar usuario",
            cuentas_eliminables,
            index=None,
            placeholder="Selecciona un usuario",
            disabled=not cuentas_eliminables,
        )
        contrasena_admin = st.text_input(
            "Contraseña de administrador",
            type="password",
        )
        confirmar_baja = st.checkbox(
            "Confirmo que deseo dar de baja esta cuenta",
            disabled=not cuentas_eliminables,
        )
        dar_de_baja = st.form_submit_button(
            "Dar de baja esta cuenta",
            type="primary",
            use_container_width=True,
            disabled=not cuentas_eliminables,
        )

    if not dar_de_baja:
        if not cuentas_eliminables:
            st.info("No hay cuentas disponibles para dar de baja.")
        return

    cuenta_admin = usuarios.get(st.session_state.get("usuario_actual"))
    cuenta_objetivo = usuarios.get(usuario_seleccionado)
    if (
        cuenta_admin is None
        or not cuenta_admin["activo"]
        or cuenta_admin["rol"] != "Administrador"
        or not verificar_contrasena(cuenta_admin, contrasena_admin)
    ):
        st.error("La contraseña de administrador no es válida.")
    elif not confirmar_baja:
        st.warning("Confirma que deseas dar de baja esta cuenta.")
    elif (
        cuenta_objetivo is None
        or not cuenta_objetivo["activo"]
        or cuenta_objetivo["rol"] == "Administrador"
        or usuario_seleccionado == st.session_state.usuario_actual
    ):
        st.error("No está permitido dar de baja esta cuenta.")
    else:
        usuarios_actualizados = {
            nombre: cuenta.copy() for nombre, cuenta in usuarios.items()
        }
        usuarios_actualizados[usuario_seleccionado]["activo"] = False
        try:
            guardar_usuarios(
                usuarios_actualizados,
                usuarios_eliminados_definitivamente=(
                    usuarios_eliminados_definitivamente
                ),
            )
        except (psycopg2.Error, ValueError) as error:
            st.error(f"No se pudo dar de baja al usuario: {error}")
        else:
            st.session_state["mensaje_gestion_usuarios"] = (
                f"La cuenta {usuario_seleccionado} fue dada de baja."
            )
            st.rerun()


try:
    usuarios, usuarios_eliminados_definitivamente = cargar_usuarios()
except (psycopg2.Error, ValueError) as error:
    st.error(str(error))
    st.stop()

VERSION_AUTENTICACION = 3

if st.session_state.get("_version_autenticacion") != VERSION_AUTENTICACION:
    st.session_state.login_completado = False
    st.session_state.admin_autorizado = False
    st.session_state.pop("_registros_sesion", None)
    st.session_state.pop("_alcance_registros_sesion", None)
    st.session_state.pop("usuario_actual", None)
    st.session_state.pop("rol_usuario", None)
    st.session_state._version_autenticacion = VERSION_AUTENTICACION
elif "login_completado" not in st.session_state:
    st.session_state.login_completado = False

if not st.session_state.login_completado:
    st.markdown(
        """
        <style>
        body:has(.st-key-login_shell) [data-testid="stHeader"],
        body:has(.st-key-login_shell) [data-testid="stSidebar"] {
            display: none !important;
        }
        body:has(.st-key-login_shell) [data-testid="stMainBlockContainer"] {
            width: 100%;
            max-width: none !important;
            min-height: 100vh;
            padding: 0 !important;
        }
        body:has(.st-key-login_shell) [data-testid="stAppViewContainer"] {
            background: #090d0a;
        }
        .st-key-login_shell {
            display: flex;
            align-items: center;
            justify-content: center;
            box-sizing: border-box;
            width: 100%;
            min-height: 100vh;
            padding: 1.5rem;
            background-color: #090d0a;
            background-image:
                linear-gradient(rgba(8, 12, 9, 0.94), rgba(8, 12, 9, 0.94)),
                repeating-linear-gradient(0deg, transparent 0 31px, rgba(166, 207, 126, 0.035) 32px),
                repeating-linear-gradient(90deg, transparent 0 31px, rgba(166, 207, 126, 0.035) 32px);
        }
        .st-key-login_shell > [data-testid="stLayoutWrapper"] {
            width: min(420px, calc(100vw - 3rem));
            max-width: 420px;
            flex: 0 1 420px;
        }
        .st-key-login_card {
            box-sizing: border-box;
            width: 100%;
            padding: 2rem;
            border: 1px solid rgba(163, 230, 53, 0.2);
            border-top: 3px solid #a3e635;
            border-radius: 8px;
            background: #121714;
            box-shadow: 0 24px 64px rgba(0, 0, 0, 0.42);
        }
        .st-key-login_card > [data-testid="stVerticalBlock"] {
            gap: 0.75rem;
        }
        .login-brand {
            display: flex;
            justify-content: center;
            margin: 0.25rem 0 1rem;
        }
        .login-brand img {
            display: block;
            width: 190px;
            max-width: 80%;
            height: auto;
        }
        .login-eyebrow {
            margin: 0;
            color: #a3e635;
            font-size: 0.72rem;
            font-weight: 700;
            letter-spacing: 0.12em;
            text-align: center;
        }
        .login-title {
            margin: 0 0 0.5rem;
            color: #f1f5ef;
            font-size: 1.55rem;
            font-weight: 650;
            text-align: center;
        }
        .st-key-login_card [data-testid="stTextInput"] label {
            color: #c5cdc6;
        }
        .st-key-login_card [data-testid="stTextInput"] input {
            min-height: 2.75rem;
            border-color: rgba(192, 209, 193, 0.2);
            background: #0a0f0c;
            color: #f1f5ef;
            -webkit-text-fill-color: #f1f5ef;
        }
        .st-key-login_card [data-testid="stAlertContainer"][role="alert"] {
            border: 1px solid rgba(248, 113, 113, 0.35);
            background: rgba(120, 40, 40, 0.24);
            color: #ffd6d2;
        }
        .st-key-login_card [data-testid="stAlertContainer"][role="alert"] p,
        .st-key-login_card [data-testid="stAlertContainer"][role="alert"] svg {
            color: #ffd6d2;
            fill: #ffd6d2;
            stroke: #ffd6d2;
        }
        .st-key-login_form button {
            min-height: 2.8rem;
            border: 1px solid #a3e635;
            border-radius: 6px;
            background: #a3e635;
            color: #10160c;
            font-weight: 700;
        }
        .st-key-login_form button:hover {
            border-color: #b8f45a;
            background: #b8f45a;
            color: #10160c;
        }
        @media (max-width: 480px) {
            .st-key-login_shell {
                padding: 1rem;
            }
            .st-key-login_card {
                padding: 1.5rem;
            }
        }
        </style>
        """,
        unsafe_allow_html=True,
    )
    with st.container(key="login_shell"):
        with st.container(key="login_card"):
            st.markdown(
                f'<div class="login-brand"><img src="data:image/png;base64,{logo_base64}" alt="El Búnker"></div>',
                unsafe_allow_html=True,
            )
            st.markdown(
                '<p class="login-eyebrow">ACCESO AL SISTEMA</p><h1 class="login-title">Iniciar sesión</h1>',
                unsafe_allow_html=True,
            )
            if not usuarios:
                st.warning(
                    "No hay cuentas cargadas en Supabase. Migra los usuarios "
                    "existentes con `python migrar_usuarios_a_supabase.py --apply`."
                )
            with st.form("login_form"):
                usuario_login = st.text_input("Usuario", key="usuario_login")
                contrasena_login = st.text_input(
                    "Contraseña", type="password", key="contrasena_login"
                )
                acceder = st.form_submit_button(
                    "Acceder", type="primary", use_container_width=True
                )
            if acceder:
                nombre_usuario_login = usuario_login.strip()
                cuenta = usuarios.get(nombre_usuario_login)
                if (
                    cuenta
                    and cuenta["activo"]
                    and verificar_contrasena(cuenta, contrasena_login)
                ):
                    ultimo_inicio_sesion = datetime.now(
                        ZONA_HORARIA_PERU
                    ).isoformat(
                        timespec="seconds"
                    )
                    try:
                        actualizado = actualizar_ultimo_inicio_sesion(
                            nombre_usuario_login,
                            ultimo_inicio_sesion,
                        )
                        if not actualizado:
                            raise ValueError(
                                "La cuenta dejó de existir durante el inicio de sesión."
                            )
                    except (psycopg2.Error, ValueError) as error:
                        st.error(f"No se pudo registrar el inicio de sesión: {error}")
                    else:
                        cuenta["ultimo_inicio_sesion"] = ultimo_inicio_sesion
                        st.session_state["_usuarios_sesion"] = usuarios
                        with st.spinner("Verificando acceso..."):
                            time.sleep(0.6)
                        st.session_state.login_completado = True
                        st.session_state.usuario_actual = nombre_usuario_login
                        st.session_state.rol_usuario = cuenta["rol"]
                        st.session_state.admin_autorizado = (
                            cuenta["rol"] == "Administrador"
                        )
                        st.rerun()
                else:
                    st.error("ACCESO DENEGADO: usuario o contraseña incorrectos.")
    st.stop()

if st.session_state.login_completado:
    cuenta_actual = usuarios.get(st.session_state.get("usuario_actual"))
    if not cuenta_actual or not cuenta_actual["activo"]:
        st.session_state.login_completado = False
        st.session_state.admin_autorizado = False
        st.session_state.pop("_registros_sesion", None)
        st.session_state.pop("_alcance_registros_sesion", None)
        st.session_state.pop("_usuarios_sesion", None)
        st.session_state.pop("_usuarios_eliminados_sesion", None)
        st.session_state.pop("usuario_actual", None)
        st.session_state.pop("rol_usuario", None)
        st.rerun()
    st.session_state.rol_usuario = cuenta_actual["rol"]
    st.session_state.admin_autorizado = (
        cuenta_actual["rol"] == "Administrador"
    )

st.markdown(
    f'<style data-style-revision="{os.stat(__file__).st_mtime_ns}">'
    + """
    @import url('https://fonts.googleapis.com/css2?family=Inter:wght@400;500;600;700;800&display=swap');
    html, body, [data-testid="stAppViewContainer"], .stApp {
        font-family: "Inter", "Segoe UI", sans-serif;
        -webkit-font-smoothing: antialiased;
        -moz-osx-font-smoothing: grayscale;
        text-rendering: optimizeLegibility;
        line-height: 1.5;
    }
    h1, h2, h3 {
        font-weight: 700;
        letter-spacing: -0.03em;
        line-height: 1.15;
    }
    p, div, span, label, button, input, select, textarea {
        letter-spacing: 0.01em;
    }
    [data-testid="stMain"] .block-container {
        padding-top: 1.5rem;
        padding-bottom: 2rem;
    }
    [data-testid="stAppViewContainer"] > div {
        display: flex;
        flex-direction: column;
        min-height: 0;
    }
    [data-testid="stMain"] {
        flex: 1 1 0% !important;
        height: auto !important;
        min-height: 0 !important;
        overflow-y: auto !important;
    }
    .stTabs [role="tablist"] button,
    .stButton > button,
    [data-testid="stTextInput"] label,
    [data-testid="stSelectbox"] label,
    [data-testid="stNumberInput"] label,
    [data-testid="stDateInput"] label {
        font-size: 0.96rem;
        line-height: 1.4;
    }
    header,
    [data-testid="stHeader"] {
        display: block !important;
        visibility: visible !important;
        position: relative !important;
        inset: auto !important;
        z-index: 10 !important;
        background: transparent !important;
        box-shadow: none !important;
    }
    body:has(#registros-theme-marker) [data-testid="stHeader"] {
        isolation: isolate;
        border-bottom: 1px solid rgba(255, 255, 255, 0.08);
        background: transparent !important;
        box-shadow: 0 4px 14px rgba(0, 0, 0, 0.16) !important;
    }
    [data-testid="stAppDeployButton"],
    [data-testid="stMainMenu"] {
        display: none !important;
        visibility: hidden !important;
    }
    [data-testid="stSidebarCollapsedControl"] {
        position: fixed !important;
        top: 0.5rem !important;
        left: 0.5rem !important;
        z-index: 10001 !important;
        display: flex !important;
        visibility: visible !important;
        opacity: 1 !important;
        pointer-events: auto !important;
    }
    [data-testid="stSidebarCollapsedControl"] button {
        min-width: 2.5rem !important;
        min-height: 2.5rem !important;
        border: 1px solid rgba(163, 230, 53, 0.5) !important;
        border-radius: 0.65rem !important;
        background: #151d16 !important;
        color: #d9f99d !important;
        cursor: pointer !important;
    }
    [data-testid="stSidebarCollapsedControl"] button:hover {
        border-color: #a3e635 !important;
        background: #202b20 !important;
    }
    .st-key-menu_usuario {
        position: fixed !important;
        top: 0.5rem;
        right: 1rem;
        z-index: 1000001 !important;
        width: 52px;
        pointer-events: auto !important;
    }
    .st-key-menu_usuario button {
        display: flex;
        align-items: center;
        justify-content: center;
        width: 48px;
        min-width: 48px;
        height: 48px;
        padding: 0;
        border: 1px solid rgba(163, 230, 53, 0.5);
        border-radius: 50%;
        background: #151d16;
        color: #d9f99d;
        cursor: pointer !important;
        pointer-events: auto !important;
    }
    .st-key-menu_usuario button:hover {
        border-color: #a3e635;
        background: #202b20;
        color: #d9f99d;
    }
    body:has(
        .st-key-menu_usuario
        [data-testid="stPopoverButton"][aria-expanded="true"]
    )
    [data-testid="stPopoverBody"] {
        z-index: 1000002 !important;
        width: 220px !important;
        padding: 0.85rem;
        border: 1px solid rgba(163, 230, 53, 0.22);
        border-radius: 8px;
        background: #121714;
        box-shadow: 0 12px 32px rgba(0, 0, 0, 0.4);
    }
    body:has(
        .st-key-menu_usuario
        [data-testid="stPopoverButton"][aria-expanded="true"]
    )
    [data-testid="stPopoverBody"] [data-testid="stVerticalBlock"] {
        gap: 0.5rem;
    }
    .profile-menu-avatar {
        display: flex;
        justify-content: center;
        margin: 0.25rem 0 0.75rem;
        color: #a3e635;
        font-size: 2rem;
        line-height: 1;
    }
    .profile-menu-avatar span {
        font-family: "Material Symbols Rounded";
        font-feature-settings: "liga";
        font-size: 2rem;
    }
    .profile-menu-label {
        margin: 0.35rem 0 0;
        color: #879188;
        font-size: 0.72rem;
        font-weight: 600;
        text-align: center;
        text-transform: uppercase;
    }
    .profile-menu-role {
        margin: 0;
        color: #f1f5ef;
        font-size: 1rem;
        font-weight: 700;
        text-align: center;
    }
    .profile-menu-user {
        margin: 0;
        color: #f1f5ef;
        font-weight: 600;
        text-align: center;
    }
    body:has(#registros-theme-marker) .block-container {
        padding-top: 10px !important;
        padding-bottom: calc(2rem + 20px) !important;
    }
    .sidebar-logo-container {
        display: flex;
        justify-content: center;
        padding: 0.25rem 0 1rem;
    }
    .sidebar-logo-container img {
        display: block;
        width: 88%;
        max-width: 176px;
        height: auto;
    }
    .main-header-logo {
        display: flex;
        justify-content: flex-end;
        margin-bottom: 0.5rem;
    }
    .main-header-logo img {
        display: block;
        width: 112px;
        max-width: 100%;
        height: auto;
    }
    .st-key-registros_hero {
        position: relative;
        isolation: isolate;
        width: 100%;
        min-height: clamp(14rem, 22vw, 19rem);
        margin-top: -6.75rem;
        margin-bottom: 1rem;
        padding: calc(1.25rem + 3rem) 1.25rem 1.25rem;
        box-sizing: border-box;
    }
    .st-key-registros_hero::before {
        content: "";
        position: absolute;
        top: 60px;
        right: -80px;
        left: -80px;
        height: 271px;
        z-index: 0;
        pointer-events: none;
        background-size: cover;
        background-position: center;
        background-repeat: no-repeat;
    }
    .st-key-registros_hero .main-header-logo {
        transform: translateY(10px);
    }
    .st-key-registros_hero > [data-testid="stLayoutWrapper"] {
        position: relative;
        z-index: 1;
    }
    .reportes-kpi-panel {
        box-sizing: border-box;
        width: 100%;
        margin: 0.75rem 0 1rem;
        padding: 1rem;
        border: 1px solid rgba(173, 190, 174, 0.18);
        border-radius: 12px;
        background: rgba(13, 19, 15, 0.72);
        backdrop-filter: blur(14px);
        box-shadow: inset 0 1px 0 rgba(255, 255, 255, 0.035),
            0 12px 32px rgba(0, 0, 0, 0.22);
    }
    .reportes-kpi-header {
        display: flex;
        align-items: center;
        justify-content: space-between;
        gap: 0.75rem;
        margin-bottom: 0.85rem;
    }
    .reportes-kpi-heading {
        color: #e6ece5;
        font-size: 0.95rem;
        font-weight: 700;
    }
    .reportes-kpi-periodo {
        flex: 0 0 auto;
        padding: 0.3rem 0.65rem;
        border: 1px solid rgba(163, 230, 53, 0.28);
        border-radius: 999px;
        background: rgba(163, 230, 53, 0.09);
        color: #d9f99d;
        font-size: 0.72rem;
        font-weight: 700;
    }
    .reportes-kpi-grid {
        display: grid;
        grid-template-columns: repeat(3, minmax(0, 1fr));
        gap: 0.7rem;
    }
    .reportes-kpi-card {
        position: relative;
        min-width: 0;
        box-sizing: border-box;
        overflow: hidden;
        padding: 1rem 0.9rem 0.85rem;
        border: 1px solid rgba(176, 190, 177, 0.15);
        border-radius: 10px;
        background: rgba(24, 31, 26, 0.7);
    }
    .reportes-kpi-acento {
        position: absolute;
        top: 0;
        left: 0;
        width: 100%;
        height: 3px;
        background: #a3e635;
    }
    .reportes-kpi-card.supervisor .reportes-kpi-acento {
        background: #93c5fd;
    }
    .reportes-kpi-card.gasto .reportes-kpi-acento {
        background: #f87171;
    }
    .reportes-kpi-card.total .reportes-kpi-acento {
        background: #fbbf64;
    }
    .reportes-kpi-card.acumulado .reportes-kpi-acento {
        background: #a3e635;
    }
    .reportes-kpi-card.finanzas-metodo-card {
        --finanzas-metodo-color: #93c5fd;
        --finanzas-metodo-color-suave: rgba(147, 197, 253, 0.28);
    }
    .reportes-kpi-card.finanzas-metodo-card .reportes-kpi-acento {
        background: var(--finanzas-metodo-color);
    }
    .reportes-kpi-card.finanzas-metodo-card::after {
        position: absolute;
        top: 0;
        right: 0;
        width: 9px;
        height: 100%;
        background: linear-gradient(
            180deg,
            var(--finanzas-metodo-color),
            var(--finanzas-metodo-color-suave)
        );
        clip-path: polygon(38% 0, 100% 0, 100% 100%, 0 100%);
        content: "";
    }
    .reportes-kpi-card.finanzas-metodo-efectivo {
        --finanzas-metodo-color: #4ade80;
        --finanzas-metodo-color-suave: rgba(74, 222, 128, 0.28);
    }
    .reportes-kpi-card.finanzas-metodo-yape {
        --finanzas-metodo-color: #c084fc;
        --finanzas-metodo-color-suave: rgba(192, 132, 252, 0.28);
    }
    .reportes-kpi-card.finanzas-metodo-plin {
        --finanzas-metodo-color: #67e8f9;
        --finanzas-metodo-color-suave: rgba(103, 232, 249, 0.28);
    }
    .reportes-kpi-card.finanzas-metodo-transferencia {
        --finanzas-metodo-color: #fb923c;
        --finanzas-metodo-color-suave: rgba(251, 146, 60, 0.28);
    }
    .reportes-kpi-card.finanzas-metodo-tarjeta {
        --finanzas-metodo-color: #2563eb;
        --finanzas-metodo-color-suave: rgba(37, 99, 235, 0.32);
    }
    .reportes-kpi-etiqueta {
        overflow: hidden;
        color: #aeb9b0;
        font-size: 0.72rem;
        font-weight: 600;
        text-overflow: ellipsis;
        white-space: nowrap;
    }
    .reportes-kpi-icono {
        margin-bottom: 0.55rem;
        color: #a3e635;
        font-family: "Material Symbols Rounded";
        font-feature-settings: "liga";
        font-size: 1.2rem;
        line-height: 1;
    }
    .reportes-kpi-card.supervisor .reportes-kpi-icono {
        color: #93c5fd;
    }
    .reportes-kpi-card.gasto .reportes-kpi-icono {
        color: #f87171;
    }
    .reportes-kpi-card.total .reportes-kpi-icono {
        color: #fbbf64;
    }
    .reportes-kpi-card.acumulado .reportes-kpi-icono {
        color: #a3e635;
    }
    .reportes-kpi-valor {
        margin-top: 0.3rem;
        color: #f4f7f2;
        font-size: 2rem;
        font-weight: 750;
        line-height: 1.1;
    }
    .reportes-kpi-unidad {
        margin-top: 0.2rem;
        color: #879188;
        font-size: 0.68rem;
    }
    .finanzas-kpi-grid {
        display: grid;
        grid-template-columns: repeat(auto-fit, minmax(190px, 1fr));
        gap: 0.7rem;
    }
    .finanzas-balance-grid {
        display: grid;
        grid-template-columns: repeat(2, minmax(0, 1fr));
        gap: 0.85rem;
        margin-bottom: 1rem;
    }
    .finanzas-balance-card {
        position: relative;
        min-width: 0;
        box-sizing: border-box;
        overflow: hidden;
        padding: 1rem 1.1rem;
        border: 1px solid rgba(176, 190, 177, 0.15);
        border-radius: 10px;
        background: rgba(24, 31, 26, 0.7);
        backdrop-filter: blur(14px);
        box-shadow: inset 0 1px 0 rgba(255, 255, 255, 0.035),
            0 8px 20px rgba(0, 0, 0, 0.16);
    }
    .finanzas-balance-card::before {
        position: absolute;
        inset: 0 0 auto;
        height: 3px;
        background: #93c5fd;
        content: "";
    }
    .finanzas-balance-card.gasto::before {
        background: #f87171;
    }
    .finanzas-balance-etiqueta {
        color: #aeb9b0;
        font-size: 0.82rem;
        font-weight: 600;
    }
    .finanzas-balance-valor {
        margin-top: 0.35rem;
        color: #f4f7f2;
        font-size: clamp(1.65rem, 3vw, 2.1rem);
        font-weight: 750;
        letter-spacing: -0.03em;
        line-height: 1.15;
    }
    .st-key-fecha_desglose_finanzas {
        max-width: 15rem;
    }
    .st-key-fecha_referencia_balance_finanzas {
        max-width: 15rem;
    }
    @media (max-width: 640px) {
        .st-key-registros_hero {
            padding: 1rem;
            padding-top: calc(1rem + 3rem);
        }
        .reportes-kpi-grid {
            grid-template-columns: repeat(2, minmax(0, 1fr));
        }
        .finanzas-kpi-grid {
            grid-template-columns: repeat(2, minmax(0, 1fr));
        }
        .finanzas-balance-grid {
            grid-template-columns: 1fr;
        }
        .reportes-kpi-panel {
            padding: 0.75rem;
        }
    }
    @media (min-width: 641px) and (max-width: 1000px) {
        .reportes-kpi-grid {
            grid-template-columns: repeat(3, minmax(0, 1fr));
        }
    }
    [data-testid="stTextInput"] label,
    [data-testid="stSelectbox"] label,
    [data-testid="stNumberInput"] label {
        color: #aaa;
        font-weight: 500;
    }
    [data-testid="stVerticalBlockBorderWrapper"] {
        background: rgba(128, 139, 150, 0.07);
        border: 1px solid rgba(128, 139, 150, 0.24);
        border-radius: 12px;
    }
    .st-key-ingresar_admin button {
        background-color: #F8D7DA !important;
        border-color: #E9B9BE !important;
        color: #842029 !important;
        -webkit-text-fill-color: #842029 !important;
    }
    .st-key-ingresar_admin button:hover {
        background-color: #F1C2C7 !important;
        border-color: #DFA4AB !important;
    }
    .st-key-abrir_agregar_personal button,
    .st-key-abrir_agregar_tipo_vehiculo button,
    .st-key-abrir_agregar_servicio button,
    .st-key-guardar_registro_btn button {
        background-color: #BFE8CE !important;
        border-color: #95D3AA !important;
        color: #174B2B !important;
        -webkit-text-fill-color: #174B2B !important;
    }
    .st-key-ingresar_admin button:hover {
        background-color: #EEB2AD !important;
        border-color: #D98E87 !important;
    }
    .st-key-quitar_tipo_button button,
    .st-key-quitar_servicio_button button,
    .st-key-dar_baja_personal_button button,
    .st-key-eliminar_registro_btn button {
        background-color: #F4C1BB !important;
        border-color: #E69A91 !important;
        color: #702B25 !important;
        -webkit-text-fill-color: #702B25 !important;
    }
    .st-key-quitar_tipo_button button:hover,
    .st-key-quitar_servicio_button button:hover,
    .st-key-dar_baja_personal_button button:hover,
    .st-key-eliminar_registro_btn button:hover {
        background-color: #EFA8A0 !important;
        border-color: #D98279 !important;
    }
    .st-key-cancelar_edicion_btn button {
        background-color: #DDF2FF !important;
        border-color: #B9DDF2 !important;
        color: #075985 !important;
        -webkit-text-fill-color: #075985 !important;
    }
    .st-key-cancelar_edicion_btn button:hover {
        background-color: #C6E8FA !important;
        border-color: #A4D5EF !important;
    }
    .st-key-abrir_agregar_personal button:hover,
    .st-key-abrir_agregar_tipo_vehiculo button:hover,
    .st-key-abrir_agregar_servicio button:hover,
    .st-key-guardar_registro_btn button:hover {
        background-color: #A8DDBA !important;
        border-color: #7FC693 !important;
    }
    [data-testid="stAlertContainer"] {
        border-radius: 10px !important;
        backdrop-filter: blur(12px);
        -webkit-backdrop-filter: blur(12px);
        box-shadow: inset 0 1px 0 rgba(255, 255, 255, 0.08),
            0 8px 24px rgba(0, 0, 0, 0.2);
    }
    [data-testid="stAlertContainer"][role="alert"] {
        background-color: rgba(120, 40, 40, 0.4) !important;
        border: 1px solid rgba(248, 113, 113, 0.3) !important;
        color: #FDE8E7 !important;
        -webkit-text-fill-color: #FDE8E7 !important;
    }
    [data-testid="stAlertContainer"][role="alert"] p,
    [data-testid="stAlertContainer"][role="alert"] svg {
        color: #FDE8E7 !important;
        fill: #FDE8E7 !important;
        stroke: #FDE8E7 !important;
    }
    .st-key-notificacion_exito [data-testid="stAlertContainer"] {
        background-color: rgba(30, 80, 50, 0.4) !important;
        border: 1px solid rgba(74, 222, 128, 0.3) !important;
        color: #E3F8EA !important;
        -webkit-text-fill-color: #E3F8EA !important;
    }
    .st-key-notificacion_exito [data-testid="stAlertContainer"] p,
    .st-key-notificacion_exito [data-testid="stAlertContainer"] svg {
        color: #E3F8EA !important;
        fill: #E3F8EA !important;
        stroke: #E3F8EA !important;
    }
    .st-key-notificacion_eliminacion [data-testid="stAlertContainer"] {
        background-color: rgba(120, 40, 40, 0.4) !important;
        border: 1px solid rgba(248, 113, 113, 0.3) !important;
        color: #FDE8E7 !important;
        -webkit-text-fill-color: #FDE8E7 !important;
    }
    .st-key-notificacion_eliminacion [data-testid="stAlertContainer"] p,
    .st-key-notificacion_eliminacion [data-testid="stAlertContainer"] svg {
        color: #FDE8E7 !important;
        fill: #FDE8E7 !important;
        stroke: #FDE8E7 !important;
    }
    [data-testid="stTextInput"] input,
    [data-testid="stNumberInput"] input,
    [data-testid="stDateInput"] input,
    [data-testid="stTimeInput"] input {
        background-color: #fff !important;
        color: #111 !important;
        -webkit-text-fill-color: #111;
    }
    [data-testid="stSelectbox"] [role="group"],
    [data-testid="stMultiSelect"] [role="group"] {
        background-color: #fff !important;
        color: #111 !important;
        border-radius: 6px;
    }
    [data-testid="stSelectbox"] [role="combobox"],
    [data-testid="stSelectbox"] input,
    [data-testid="stMultiSelect"] [role="combobox"],
    [data-testid="stMultiSelect"] input {
        background-color: #fff !important;
        color: #111 !important;
        -webkit-text-fill-color: #111;
    }
    [data-testid="stSelectbox"] [role="group"] button,
    [data-testid="stMultiSelect"] [role="group"] button {
        background-color: #B3262D !important;
        border-color: #B3262D !important;
        color: #fff !important;
    }
    [data-testid="stSelectbox"] [role="group"] button svg,
    [data-testid="stMultiSelect"] [role="group"] button svg {
        color: #fff !important;
        stroke: #fff !important;
    }
    [data-testid="stSelectbox"] [role="group"] button:hover,
    [data-testid="stMultiSelect"] [role="group"] button:hover {
        background-color: #92222A !important;
        border-color: #92222A !important;
    }
    [data-testid="stTextInput"] input::placeholder,
    [data-testid="stNumberInput"] input::placeholder,
    [data-testid="stDateInput"] input::placeholder,
    [data-testid="stTimeInput"] input::placeholder {
        color: #555 !important;
        -webkit-text-fill-color: #555;
    }
    [data-testid="stPopoverBody"]:has([class*="_mes_calendario"]) {
        z-index: 1000 !important;
        width: min(340px, calc(100vw - 24px)) !important;
        min-width: 0 !important;
        max-height: calc(100vh - 24px) !important;
        overflow-y: auto !important;
        overscroll-behavior: contain;
        box-sizing: border-box !important;
    }
    @media (max-width: 640px) {
        [data-testid="stPopoverBody"]:has([class*="_mes_calendario"])
        [data-testid="stHorizontalBlock"]:not(:has([data-testid="stSelectbox"])) {
            display: grid !important;
            grid-template-columns: repeat(7, minmax(0, 1fr)) !important;
            align-items: stretch !important;
        }
        [data-testid="stPopoverBody"]:has([class*="_mes_calendario"])
        [data-testid="stHorizontalBlock"]:not(:has([data-testid="stSelectbox"]))
        > [data-testid="column"] {
            width: auto !important;
            min-width: 0 !important;
            max-width: none !important;
        }
    }
    [class*="_popover_calendario"] [data-testid="stPopoverButton"] {
        min-height: 40px !important;
        padding: 0.45rem 0.75rem !important;
        border: 1px solid rgba(173, 190, 174, 0.35) !important;
        border-radius: 8px !important;
        background: #151d16 !important;
        color: #f1f5ef !important;
        visibility: visible !important;
    }
    [class*="_popover_calendario"] [data-testid="stPopoverButton"]:hover {
        border-color: #a3e635 !important;
        background: #202b20 !important;
    }
    [data-testid="stPopoverBody"]:has([class*="_dia_"])
    [class*="_dia_"] [data-testid^="stBaseButton-"] {
        box-sizing: border-box !important;
        width: 100% !important;
        min-width: 0 !important;
        height: 32px !important;
        min-height: 32px !important;
        padding: 0 !important;
        font-size: 0.82rem !important;
        line-height: 1 !important;
    }
    [data-testid="stPopoverBody"]:has([class*="_mes_calendario"])
    [data-testid="stSelectbox"] [role="group"],
    [data-testid="stPopoverBody"]:has([class*="_mes_calendario"])
    [data-testid="stSelectbox"] [role="combobox"] {
        background-color: #202922 !important;
        color: #f4f7f2 !important;
        -webkit-text-fill-color: #f4f7f2 !important;
    }
    [data-testid$="VirtualDropdown"]:not(
        :has([role="listbox"][aria-label="Mes"])
    ):not(:has([role="listbox"][aria-label="Año"])) {
        border: 1px solid #d1d5db !important;
        border-radius: 8px !important;
        background-color: #fff !important;
        color: #111 !important;
        box-shadow: 0 8px 24px rgba(0, 0, 0, 0.2) !important;
    }
    [data-testid$="VirtualDropdown"]:not(
        :has([role="listbox"][aria-label="Mes"])
    ):not(:has([role="listbox"][aria-label="Año"]))
    [role="listbox"] {
        background-color: #fff !important;
        color: #111 !important;
    }
    [data-testid$="VirtualDropdown"]:not(
        :has([role="listbox"][aria-label="Mes"])
    ):not(:has([role="listbox"][aria-label="Año"]))
    [role="option"],
    [data-testid$="VirtualDropdown"]:not(
        :has([role="listbox"][aria-label="Mes"])
    ):not(:has([role="listbox"][aria-label="Año"]))
    [role="option"] [data-item-hl] {
        background-color: transparent !important;
        color: #111 !important;
        -webkit-text-fill-color: #111 !important;
    }
    [data-testid$="VirtualDropdown"]:not(
        :has([role="listbox"][aria-label="Mes"])
    ):not(:has([role="listbox"][aria-label="Año"]))
    [role="option"][aria-selected="true"] [data-item-hl],
    [data-testid$="VirtualDropdown"]:not(
        :has([role="listbox"][aria-label="Mes"])
    ):not(:has([role="listbox"][aria-label="Año"]))
    [role="option"][data-focused="true"] [data-item-hl] {
        border-radius: 6px;
        background-color: #e5e7eb !important;
        color: #111 !important;
    }
    .st-key-descargar_registros_excel button,
    [data-testid="stDownloadButton"] button {
        background-color: #198754 !important;
        border-color: #198754 !important;
        color: #fff !important;
    }
    .st-key-descargar_registros_excel button:hover,
    [data-testid="stDownloadButton"] button:hover {
        background-color: #146c43 !important;
        border-color: #146c43 !important;
    }
    body:has(#registros-theme-marker) [data-testid="stDataFrame"] canvas[style*="height: 36px"] {
        filter: sepia(1) saturate(6) hue-rotate(314deg) brightness(2.8);
    }
    body:has(#registros-theme-marker)
    [class*="st-key-tabla_admin_"] [data-testid="stDataFrame"],
    body:has(#registros-theme-marker)
    [class*="st-key-tabla_registros_supervisor"] [data-testid="stDataFrame"] {
        position: relative !important;
    }
    body:has(#registros-theme-marker)
    [class*="st-key-tabla_admin_"] [data-testid="stDataFrame"]::after,
    body:has(#registros-theme-marker)
    [class*="st-key-tabla_registros_supervisor"] [data-testid="stDataFrame"]::after {
        content: "";
        position: absolute;
        inset: 0 0 auto;
        height: 36px;
        z-index: 10;
        pointer-events: auto;
        background: transparent;
    }
    """ + "</style>",
    unsafe_allow_html=True,
)
st.markdown(
    f"""
    <style>
    [data-testid="stSidebar"] {{
        background-image:
            linear-gradient(rgba(15, 20, 25, 0.85), rgba(15, 20, 25, 0.85)),
            url("data:image/png;base64,{fondo_sidebar_base64}");
        background-size: cover;
        background-position: center;
        background-repeat: no-repeat;
    }}
    </style>
    """,
    unsafe_allow_html=True,
)
st.markdown(
    f"""
    <style>
    body:has(#registros-theme-marker) [data-testid="stHeader"] {{
        position: fixed !important;
        top: 0 !important;
        left: 0 !important;
        right: 0 !important;
        width: auto !important;
        height: 38px !important;
        min-height: 38px !important;
        z-index: 10000 !important;
        overflow: visible !important;
        background: transparent !important;
    }}
    body:has(#registros-theme-marker) .st-key-menu_usuario {{
        top: 3px !important;
        right: 8px !important;
        width: 32px !important;
        z-index: 10001 !important;
    }}
    body:has(#registros-theme-marker) .st-key-menu_usuario button {{
        width: 32px !important;
        min-width: 32px !important;
        max-width: 32px !important;
        height: 32px !important;
        min-height: 32px !important;
        max-height: 32px !important;
        padding: 0 !important;
    }}
    body:has(#registros-theme-marker) [data-testid="stHeader"]::after {{
        content: "";
        position: absolute;
        inset: 0 0 auto;
        height: 38px;
        z-index: 0;
        pointer-events: none;
        border-bottom: 1px solid rgba(255, 255, 255, 0.14);
        background: rgba(15, 20, 25, 0.08);
        backdrop-filter: blur(5px);
        -webkit-backdrop-filter: blur(5px);
        box-shadow: 0 4px 14px rgba(0, 0, 0, 0.2);
    }}
    .st-key-registros_hero::before {{
        background-image: url("data:image/png;base64,{banner_base64}");
    }}
    body:has(#registros-theme-marker) .st-key-registros_hero > [data-testid="stLayoutWrapper"] {{
        transform: translateY(38px);
    }}
    </style>
    """,
    unsafe_allow_html=True,
)

DIRECTORIO_APP = os.path.dirname(os.path.abspath(__file__))
ARCHIVO_LAVADORES = "lavadores.txt"
ARCHIVO_VEHICULOS = "vehiculos.txt"
ARCHIVO_SERVICIOS = "servicios.txt"
COLUMNAS_CLIENTE = [
    "RUC/DNI",
    "Razón Social/Nombre",
    "Número Cliente",
    "Correo Cliente",
]
NOMBRE_COLUMNA_LOCAL = "Local"
USUARIOS_SUPERVISORES = sorted(
    nombre
    for nombre, cuenta in usuarios.items()
    if cuenta["activo"] and cuenta["rol"] == "Supervisor"
)
COLUMNAS_VALIDAS = [
    "ID", "Tipo", "Placa", "Servicio", "Pago", "Monto", "Fecha",
    "Lavador 1", "Lavador 2", "Lavador 3", "Motivo Gasto",
    "Precio Gasto", "Método Gasto",
    *COLUMNAS_CLIENTE, "Registrado por",
]
COLUMNAS_EXPORTACION = [
    ("Tipo", "TIPO"),
    ("Placa", "PLACA"),
    ("Servicio", "SERVICIO"),
    ("Pago", "PAGO"),
    ("Monto", "MONTO"),
    ("Fecha", "FECHA"),
    ("Hora", "HORA"),
    ("Lavador 1", "LAVADOR 1"),
    ("Lavador 2", "LAVADOR 2"),
    ("Lavador 3", "LAVADOR 3"),
    ("Motivo Gasto", "MOTIVO GASTO"),
    ("Precio Gasto", "PRECIO GASTO"),
    ("Método Gasto", "MÉTODO GASTO"),
    ("RUC/DNI", "RUC/DNI CLIENTE"),
    ("Razón Social/Nombre", "RAZÓN SOCIAL/NOMBRE"),
    ("Número Cliente", "NÚMERO CLIENTE"),
    ("Correo Cliente", "CORREO CLIENTE"),
]

def cargar_catalogo(archivo, valores_iniciales):
    if os.path.exists(archivo):
        with open(archivo, "r", encoding="utf-8") as catalogo:
            valores = [linea.strip() for linea in catalogo if linea.strip()]
        return valores
    return valores_iniciales.copy()

def guardar_catalogo(archivo, valores):
    with open(archivo, "w", encoding="utf-8") as catalogo:
        catalogo.write("\n".join(valores) + "\n")

def consultar_dataframe(fecha=None, supervisor=None, mes=None):
    df = _obtener_registros_sesion()
    if fecha is not None and not df.empty:
        fechas = df["Fecha"].map(parsear_fecha)
        df = df.loc[
            fechas.map(
                lambda valor: pd.notna(valor) and valor.date() == fecha
            )
        ]
    if mes is not None and not df.empty:
        fechas = df["Fecha"].map(parsear_fecha)
        df = df.loc[
            fechas.map(
                lambda valor: pd.notna(valor)
                and valor.year == mes.year
                and valor.month == mes.month
            )
        ]
    if supervisor is not None:
        responsables = (
            df["Registrado por"].fillna("").astype(str).str.strip().str.casefold()
        )
        df = df.loc[responsables == str(supervisor).strip().casefold()]
    return df.reindex(columns=COLUMNAS_VALIDAS, fill_value="").copy()


def _obtener_registros_sesion(forzar_recarga=False):
    nombre_usuario = st.session_state.get("usuario_actual")
    cuenta = usuarios.get(nombre_usuario, {})
    rol = cuenta.get("rol", st.session_state.get("rol_usuario", ""))
    alcance = (nombre_usuario, rol)
    if (
        forzar_recarga
        or st.session_state.get("_alcance_registros_sesion") != alcance
        or "_registros_sesion" not in st.session_state
    ):
        supervisor = nombre_usuario if rol == "Supervisor" else None
        filas = consultar_registros(supervisor=supervisor)
        df = pd.DataFrame.from_records(
            filas,
            columns=COLUMNAS_REGISTRO.values(),
        )
        df = df.rename(
            columns={
                nombre_sql: nombre_app
                for nombre_app, nombre_sql in COLUMNAS_REGISTRO.items()
            }
        )
        st.session_state["_registros_sesion"] = df.reindex(
            columns=COLUMNAS_VALIDAS, fill_value=""
        )
        st.session_state["_alcance_registros_sesion"] = alcance
    return st.session_state["_registros_sesion"].copy()


def refrescar_registros_sesion():
    if st.session_state.get("login_completado"):
        _obtener_registros_sesion(forzar_recarga=True)


def guardar_eliminacion_supervisor(
    usuarios_actualizados,
    nombre_supervisor,
    eliminados_actualizados,
):
    with conectar() as conexion:
        eliminar_registros_supervisor_en_conexion(
            conexion, nombre_supervisor
        )
        guardar_usuarios_en_conexion(
            conexion,
            usuarios_actualizados,
            eliminados_actualizados,
        )
    _actualizar_usuarios_en_sesion(
        usuarios_actualizados, eliminados_actualizados
    )
    refrescar_registros_sesion()

def limpiar_monto(valor):
    try:
        val_str = str(valor).replace("S/", "").strip()
        if val_str in ["", "-", "nan", "None"]:
            return 0.0
        return float(val_str)
    except:
        return 0.0

def estilo_metodo_pago(valor):
    estilos = {
        "efectivo": "background-color: #DDF3E4; color: #245B35; font-weight: 600; text-align: center;",
        "tarjeta": "background-color: #DCEBFA; color: #1E4E79; font-weight: 600; text-align: center;",
        "yape": "background-color: #EADFF5; color: #593477; font-weight: 600; text-align: center;",
        "yape/plin": "background-color: #EADFF5; color: #593477; font-weight: 600; text-align: center;",
        "plin": "background-color: #D9F1F8; color: #155E75; font-weight: 600; text-align: center;",
        "transferencia": "background-color: #FFE2C2; color: #9A3412; font-weight: 600; text-align: center;",
    }
    return estilos.get(str(valor).strip().casefold())

def estilo_fila_tabla(fila, filas_seleccionadas):
    estilo_base = "background-color: #FFFFFF; color: #111111;"
    if fila.name in filas_seleccionadas:
        estilo_base = "background-color: rgba(179, 38, 45, 0.72); color: #FFFFFF;"
    estilos_celdas = [estilo_base] * len(fila)

    estilo_pago = estilo_metodo_pago(fila["Pago"])
    if estilo_pago:
        estilos_celdas[fila.index.get_loc("Pago")] = estilo_pago
        if "Monto" in fila.index:
            estilos_celdas[fila.index.get_loc("Monto")] = estilo_pago

    if "Método Gasto" in fila.index:
        estilo_gasto = estilo_metodo_pago(fila["Método Gasto"])
        if estilo_gasto:
            estilos_celdas[fila.index.get_loc("Método Gasto")] = estilo_gasto

    if "Registrado por" in fila.index:
        rol = str(fila["Registrado por"]).strip()
        if rol == "Administrador":
            estilo_rol = (
                "background-color: #DDF3E4; color: #245B35; "
                "font-weight: 600; text-align: center; border-radius: 999px; "
                "border: 1px solid #B9DFC5;"
            )
        elif (
            rol.startswith(("Lavador", "Supervisor"))
            or normalizar_responsable(rol) in USUARIOS_SUPERVISORES
        ):
            estilo_rol = (
                "background-color: #DCEBFA; color: #1E4E79; "
                "font-weight: 600; text-align: center; border-radius: 999px; "
                "border: 1px solid #B9D5F0;"
            )
        else:
            estilo_rol = "color: #667085; text-align: center;"
        estilos_celdas[fila.index.get_loc("Registrado por")] = estilo_rol

    return estilos_celdas

def crear_tabla_estilizada(
    df,
    filas_seleccionadas=None,
    incluir_usuario=False,
):
    columnas_ocultas = ["ID"]
    if not incluir_usuario:
        columnas_ocultas.append("Registrado por")
    df_visual = (
        df.drop(columns=columnas_ocultas, errors="ignore")
        .fillna("-")
        .replace(r"^\s*$", "-", regex=True)
        .astype(str)
        .reset_index(drop=True)
    )
    if incluir_usuario and "Registrado por" in df_visual.columns:
        columnas = ["Registrado por"] + [
            columna for columna in df_visual.columns if columna != "Registrado por"
        ]
        df_visual = df_visual[columnas]
    if "Fecha" in df_visual.columns:
        fechas_separadas = df_visual["Fecha"].map(separar_fecha_hora_registro)
        posicion_hora = df_visual.columns.get_loc("Fecha") + 1
        df_visual["Fecha"] = fechas_separadas.map(lambda valor: valor[0])
        df_visual.insert(
            posicion_hora,
            "Hora",
            fechas_separadas.map(lambda valor: valor[1]),
        )
    filas_seleccionadas = filas_seleccionadas or set()
    return df_visual.style.apply(
        lambda fila: estilo_fila_tabla(fila, filas_seleccionadas), axis=1
    )


def obtener_columnas_visibles_registros(df):
    columnas = [
        columna
        for columna in df.columns
        if columna not in {"ID", "Registrado por"}
    ]
    if "Fecha" in columnas:
        columnas.insert(columnas.index("Fecha") + 1, "Hora")
    return columnas


def guardar_seleccion_tabla():
    clave_tabla = st.session_state.get("tabla_admin_key_activa", "tabla_admin")
    estado_tabla = st.session_state.get(clave_tabla)
    seleccion = estado_tabla.get("selection", {}) if estado_tabla else {}
    st.session_state[f"filas_{clave_tabla}"] = set(seleccion.get("rows", []))

def parsear_fecha(valor):
    if pd.isna(valor):
        return pd.NaT

    texto = str(valor).strip()
    formatos = [
        "%Y/%m/%d %I:%M:%S %p", "%Y/%m/%d %I:%M %p",
        "%Y/%m/%d %H:%M:%S", "%Y/%m/%d %H:%M", "%Y/%m/%d",
        "%d/%m/%Y %I:%M:%S %p", "%d/%m/%Y %I:%M %p",
        "%d/%m/%Y %H:%M:%S", "%d/%m/%Y %H:%M", "%d/%m/%Y",
        "%Y-%m-%d %I:%M:%S %p", "%Y-%m-%d %I:%M %p",
        "%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%Y-%m-%d",
        "%d-%m-%Y %I:%M:%S %p", "%d-%m-%Y %I:%M %p",
        "%d-%m-%Y %H:%M:%S", "%d-%m-%Y %H:%M", "%d-%m-%Y",
    ]
    for formato in formatos:
        try:
            return datetime.strptime(texto, formato)
        except ValueError:
            continue
    return pd.NaT


def separar_fecha_hora_registro(valor):
    texto = str(valor).strip()
    fecha = parsear_fecha(texto)
    if pd.isna(fecha):
        return texto, "-"
    hora = (
        fecha.strftime("%I:%M:%S %p")
        if re.search(r"\d{1,2}:\d{2}", texto)
        else "-"
    )
    return fecha.strftime("%Y-%m-%d"), hora


def filtrar_registros_por_fecha(df, fecha):
    if df.empty:
        return df.copy()

    fechas_normalizadas = df["Fecha"].map(parsear_fecha)
    coincide_fecha = fechas_normalizadas.map(
        lambda valor: pd.notna(valor) and valor.date() == fecha
    )
    return df.loc[coincide_fecha].copy()

def normalizar_responsable(valor):
    if pd.isna(valor):
        return "No disponible"
    texto = str(valor).strip()
    compacto = re.sub(r"[\s_-]+", "", texto).casefold()
    cuenta = next(
        (
            (nombre, detalle)
            for nombre, detalle in usuarios.items()
            if nombre.casefold() == texto.casefold()
        ),
        None,
    )
    if cuenta:
        nombre, detalle = cuenta
        return nombre if detalle["rol"] == "Supervisor" else "Administrador"
    if compacto in {"admin", "administrador"}:
        return "Administrador"
    coincidencia = re.fullmatch(r"supervisor([1-4])", compacto)
    if coincidencia:
        nombre_legacy = f"supervisor{coincidencia.group(1)}"
        if nombre_legacy in usuarios:
            return nombre_legacy
    if re.fullmatch(r"lavador[1-4]", compacto):
        return "Supervisor sin identificar"
    if compacto == "supervisor":
        return "Supervisor sin identificar"
    return texto or "No disponible"


def obtener_responsables_reporte(df):
    responsables = ["Administrador", *USUARIOS_SUPERVISORES]
    eliminados_normalizados = {
        re.sub(r"[\s_-]+", "", nombre).casefold()
        for nombre in usuarios_eliminados_definitivamente
    }
    if df.empty:
        return responsables

    historicos = df["Registrado por"].map(normalizar_responsable)
    for responsable in historicos:
        responsable_normalizado = re.sub(
            r"[\s_-]+", "", responsable
        ).casefold()
        if (
            responsable not in responsables
            and responsable_normalizado not in eliminados_normalizados
            and responsable not in {"No disponible", "Supervisor sin identificar"}
        ):
            responsables.append(responsable)
    if "Supervisor sin identificar" in historicos.values:
        responsables.append("Supervisor sin identificar")
    return responsables


def filtrar_registros_por_mes(df, fecha):
    if df.empty:
        return df.copy()

    fechas_normalizadas = df["Fecha"].map(parsear_fecha)
    coincide_mes = fechas_normalizadas.map(
        lambda valor: pd.notna(valor)
        and valor.year == fecha.year
        and valor.month == fecha.month
    )
    registros_mes = df.loc[coincide_mes].copy()
    registros_mes["_fecha_orden"] = fechas_normalizadas.loc[coincide_mes]
    return (
        registros_mes.sort_values("_fecha_orden", ascending=True)
        .drop(columns="_fecha_orden")
    )

def filtrar_registros_del_mes(df, ahora=None):
    ahora = ahora or datetime.now(ZONA_HORARIA_PERU)
    if df.empty:
        return df.copy()

    fechas = df["Fecha"].map(parsear_fecha)
    coincide_mes = fechas.map(
        lambda fecha: pd.notna(fecha)
        and fecha.year == ahora.year
        and fecha.month == ahora.month
        and fecha.date() <= ahora.date()
    )
    registros_mes = df.loc[coincide_mes].copy()
    registros_mes["_fecha_orden"] = fechas.loc[coincide_mes]
    return (
        registros_mes.sort_values("_fecha_orden", ascending=True)
        .drop(columns="_fecha_orden")
    )

def calcular_resumen_financiero(df, ahora=None):
    if df.empty:
        return 0.0, 0.0, 0.0, 0.0

    df_metricas = df.copy()
    df_metricas["Fecha_Normalizada"] = df_metricas["Fecha"].map(parsear_fecha)
    df_metricas["Monto_Num"] = df_metricas["Monto"].map(limpiar_monto)
    df_metricas["Precio_Gasto_Num"] = df_metricas["Precio Gasto"].map(limpiar_monto).abs()

    tipo_gasto = df_metricas["Tipo"].fillna("").astype(str).str.contains(
        r"\b(?:gasto|egreso|expense)\b", case=False, regex=True
    )
    df_metricas["Ingreso_Num"] = df_metricas["Monto_Num"].where(~tipo_gasto, 0.0)
    df_metricas["Gasto_Num"] = df_metricas["Precio_Gasto_Num"].where(
        df_metricas["Precio_Gasto_Num"] > 0,
        df_metricas["Monto_Num"].abs().where(tipo_gasto, 0.0),
    )

    ahora = ahora or datetime.now(ZONA_HORARIA_PERU)
    fechas = df_metricas["Fecha_Normalizada"]
    es_hoy = fechas.dt.date == ahora.date()
    es_mes_actual = (fechas.dt.year == ahora.year) & (fechas.dt.month == ahora.month)

    ganancia_dia = df_metricas.loc[es_hoy, "Ingreso_Num"].sum()
    gasto_dia = df_metricas.loc[es_hoy, "Gasto_Num"].sum()
    ganancia_mes = df_metricas.loc[es_mes_actual, "Ingreso_Num"].sum()
    gasto_mes = df_metricas.loc[es_mes_actual, "Gasto_Num"].sum()
    return ganancia_dia, gasto_dia, ganancia_mes, gasto_mes


def desglosar_metodos_financieros(df):
    metodos_base = ["Efectivo", "Yape", "Plin", "Tarjeta", "Transferencia"]
    if df.empty:
        return {metodo: 0.0 for metodo in metodos_base}, {
            metodo: 0.0 for metodo in metodos_base
        }

    df_metricas = df.copy()
    es_gasto = df_metricas["Tipo"].fillna("").astype(str).str.contains(
        r"\b(?:gasto|egreso|expense)\b", case=False, regex=True
    )
    monto_ingreso = df_metricas["Monto"].map(limpiar_monto)
    monto_gasto = df_metricas["Precio Gasto"].map(limpiar_monto).abs()
    monto_gasto = monto_gasto.where(
        monto_gasto > 0,
        monto_ingreso.abs().where(es_gasto, 0.0),
    )

    def agrupar_por_metodo(filas, columna, montos):
        metodos = filas[columna].fillna("").astype(str).str.strip()
        metodos = metodos.mask(
            metodos.str.casefold().isin({"", "-", "nan", "none"}),
            "Sin método especificado",
        )
        nombres_metodo = {
            "efectivo": "Efectivo",
            "yape": "Yape",
            "plin": "Plin",
            "tarjeta": "Tarjeta",
            "transferencia": "Transferencia",
            "yape/plin": "Yape/Plin (histórico)",
        }
        metodos = metodos.map(
            lambda metodo: nombres_metodo.get(metodo.casefold(), metodo)
        )
        return (
            montos.loc[filas.index]
            .groupby(metodos)
            .sum()
            .to_dict()
        )

    ingresos = agrupar_por_metodo(
        df_metricas.loc[~es_gasto],
        "Pago",
        monto_ingreso.loc[~es_gasto],
    )
    filas_con_gasto = monto_gasto > 0
    gastos = agrupar_por_metodo(
        df_metricas.loc[filas_con_gasto],
        "Método Gasto",
        monto_gasto.loc[filas_con_gasto],
    )

    for metodo in metodos_base:
        ingresos.setdefault(metodo, 0.0)
        gastos.setdefault(metodo, 0.0)

    return ingresos, gastos


def calcular_ingresos_por_responsable(df, fecha, periodo, hasta_fecha=None):
    orden_responsables = obtener_responsables_reporte(df)[1:]
    if df.empty:
        return pd.Series(0.0, index=orden_responsables, dtype=float)

    df_metricas = df.copy()
    df_metricas["Fecha_Normalizada"] = df_metricas["Fecha"].map(parsear_fecha)
    df_metricas["Monto_Num"] = df_metricas["Monto"].map(limpiar_monto)
    tipo_gasto = df_metricas["Tipo"].fillna("").astype(str).str.contains(
        r"\b(?:gasto|egreso|expense)\b", case=False, regex=True
    )
    df_metricas["Ingreso_Num"] = df_metricas["Monto_Num"].where(~tipo_gasto, 0.0)
    df_metricas["Responsable"] = df_metricas["Registrado por"].map(
        normalizar_responsable
    )

    fechas = df_metricas["Fecha_Normalizada"]
    if periodo == "Día":
        coincide_periodo = fechas.dt.date == fecha
    else:
        coincide_periodo = (fechas.dt.year == fecha.year) & (
            fechas.dt.month == fecha.month
        )
        if hasta_fecha is not None:
            coincide_periodo &= fechas.dt.date <= hasta_fecha

    ingresos = df_metricas.loc[coincide_periodo].groupby("Responsable")[
        "Ingreso_Num"
    ].sum()
    return ingresos.reindex(orden_responsables, fill_value=0.0)

def generar_excel_registros(df):
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    df_exportacion = df.reindex(
        columns=[columna for columna, _ in COLUMNAS_EXPORTACION],
        fill_value="-",
    ).copy()
    fechas_separadas = df_exportacion["Fecha"].map(separar_fecha_hora_registro)
    df_exportacion["Fecha"] = fechas_separadas.map(lambda valor: valor[0])
    df_exportacion["Hora"] = fechas_separadas.map(lambda valor: valor[1])
    df_exportacion.columns = [titulo for _, titulo in COLUMNAS_EXPORTACION]

    for columna in ("MONTO", "PRECIO GASTO"):
        df_exportacion[columna] = df_exportacion[columna].map(
            lambda valor: "-"
            if str(valor).strip().casefold() in {"", "-", "nan", "none"}
            else limpiar_monto(valor)
        )

    archivo = BytesIO()
    with pd.ExcelWriter(archivo, engine="openpyxl") as writer:
        df_exportacion.to_excel(writer, index=False, sheet_name="Registros")
        hoja = writer.sheets["Registros"]
        formato_moneda = '"S/ " #,##0.00;"S/ " -#,##0.00;"S/ " 0.00'
        borde_fino = Side(style="thin", color="FFB7C9D6")
        borde_celda = Border(
            left=borde_fino,
            right=borde_fino,
            top=borde_fino,
            bottom=borde_fino,
        )
        relleno_cabecera = PatternFill(fill_type="solid", fgColor="FF000000")
        fuente_cabecera = Font(color="FFFFFFFF", bold=True)

        for celda in hoja[1]:
            celda.fill = relleno_cabecera
            celda.font = fuente_cabecera
            celda.alignment = Alignment(horizontal="center", vertical="center")

        for fila in hoja.iter_rows():
            for celda in fila:
                celda.border = borde_celda
                if celda.row > 1:
                    celda.alignment = Alignment(vertical="center")

        columnas_moneda = {
            hoja.cell(row=1, column=numero_columna).value: numero_columna
            for numero_columna in range(1, hoja.max_column + 1)
            if hoja.cell(row=1, column=numero_columna).value
            in {"MONTO", "PRECIO GASTO"}
        }
        for numero_columna in columnas_moneda.values():
            for fila in range(2, hoja.max_row + 1):
                celda = hoja.cell(row=fila, column=numero_columna)
                if isinstance(celda.value, (int, float)):
                    celda.number_format = formato_moneda

        for celdas_columna in hoja.columns:
            letra_columna = celdas_columna[0].column_letter
            numero_columna = celdas_columna[0].column
            longitud_maxima = 0
            for celda in celdas_columna:
                if celda.value is None:
                    continue
                if (
                    numero_columna in columnas_moneda.values()
                    and isinstance(celda.value, (int, float))
                ):
                    texto = f"S/ {celda.value:,.2f}"
                else:
                    texto = str(celda.value)
                longitud_maxima = max(longitud_maxima, len(texto))
            hoja.column_dimensions[letra_columna].width = max(longitud_maxima + 2, 12)

        hoja.row_dimensions[1].height = 24
        hoja.freeze_panes = "A2"
        hoja.auto_filter.ref = hoja.dimensions

    return archivo.getvalue()


def validar_campos_gasto(precio_gasto, metodo_gasto):
    precio_ingresado = str(precio_gasto).strip() not in {"", "-"}
    metodo_seleccionado = str(metodo_gasto).strip() not in {"", "-"}
    errores = []
    if precio_ingresado and not metodo_seleccionado:
        errores.append(
            "No se puede registrar un precio de gasto sin seleccionar el método de gasto."
        )
    if metodo_seleccionado and not precio_ingresado:
        errores.append(
            "No se puede registrar un método de gasto sin ingresar el precio del gasto."
        )
    return errores


# --- INICIALIZAR ESQUEMA DE POSTGRESQL UNA VEZ POR SESIÓN ---
if not st.session_state.get("_esquema_postgresql_inicializado", False):
    try:
        inicializar_base_datos()
    except (OSError, ValueError, psycopg2.Error) as error:
        st.error(f"No se pudieron cargar los registros: {error}")
        st.stop()
    st.session_state["_esquema_postgresql_inicializado"] = True

def notificar_exito(mensaje, tipo="exito", destino=None):
    st.session_state.notificacion_exito = {
        "mensaje": mensaje,
        "tipo": tipo,
        "destino": destino,
        "creado_en": time.monotonic(),
    }

def mostrar_notificacion_exito(destino=None):
    notificacion = st.session_state.get("notificacion_exito")
    if notificacion:
        if isinstance(notificacion, str):
            notificacion = {"mensaje": notificacion, "tipo": "exito"}
        if notificacion.get("destino") != destino:
            return
        if destino is not None:
            if time.monotonic() - notificacion.get("creado_en", 0) >= 10:
                st.session_state.pop("notificacion_exito", None)
                return
        else:
            st.session_state.pop("notificacion_exito", None)
        tipo = notificacion.get("tipo", "exito")
        if tipo not in {"exito", "eliminacion"}:
            tipo = "exito"
        with st.container(key=f"notificacion_{tipo}"):
            st.success(notificacion["mensaje"])

# Catálogos persistentes compartidos por el panel y el formulario
lista_lavadores = cargar_catalogo(
    ARCHIVO_LAVADORES,
    [
        "Ismael",
        "Eymar",
        "Antony",
        "Anderson",
        "Alvaro",
        "Jeferson",
        "Zara",
        "Dylan",
        "Wender",
        "Eduardo",
        "Javier",
        "Eglimar",
        "Gabriel",
        "Jose",
        "Jose elis",
        "Oscar",
        "Renzo",
        "Renzo V.",
        "Sergio",
        "Kelvin",
        "Merwin",
        "William",
        "Pedro",
        "Felix",
        "Marco",
    ],
)
lista_tipos = cargar_catalogo(
    ARCHIVO_VEHICULOS, ["Auto", "Moto", "Motaxi", "Camioneta", "Bicicleta"]
)
lista_servicios = cargar_catalogo(
    ARCHIVO_SERVICIOS,
    [
        "Lavado Simple",
        "Lavado Completo",
        "Encerado",
        "Lavado De Salon",
        "L. Basico+Pulido",
        "L. Basico+Motor",
        "L. Basico+Chasis,Motor,Cera",
        "Tratamiento de cerámico",
    ],
)
if not lista_tipos:
    lista_tipos = ["Auto"]
if not lista_servicios:
    lista_servicios = ["Lavado Simple"]


@st.dialog("Agregar elemento")
def mostrar_formulario_catalogo(tipo_catalogo):
    catalogos = {
        "personal": (
            "Nuevo nombre",
            "Registrar Lavador",
            ARCHIVO_LAVADORES,
            lista_lavadores,
            "Lavador",
        ),
        "vehiculo": (
            "Nuevo tipo de vehículo",
            "Añadir tipo de vehículo",
            ARCHIVO_VEHICULOS,
            lista_tipos,
            "Tipo de vehículo",
        ),
        "servicio": (
            "Nuevo servicio",
            "Añadir servicio",
            ARCHIVO_SERVICIOS,
            lista_servicios,
            "Servicio",
        ),
    }
    etiqueta_campo, etiqueta_boton, archivo, elementos, nombre_elemento = (
        catalogos[tipo_catalogo]
    )
    with st.form(f"form_agregar_{tipo_catalogo}", clear_on_submit=True):
        nuevo_elemento = st.text_input(etiqueta_campo)
        guardar_elemento = st.form_submit_button(
            etiqueta_boton,
            type="primary",
            use_container_width=True,
        )
    if not guardar_elemento:
        return

    nombre_limpio = nuevo_elemento.strip()
    if not nombre_limpio:
        st.warning(f"Escribe {nombre_elemento.lower()} para añadir.")
    elif any(
        nombre_limpio.casefold() == existente.casefold()
        for existente in elementos
    ):
        st.warning(f"Este {nombre_elemento.lower()} ya está registrado.")
    else:
        guardar_catalogo(archivo, [*elementos, nombre_limpio])
        notificar_exito(
            f"{nombre_elemento} {nombre_limpio} añadido correctamente.",
            destino=tipo_catalogo,
        )
        st.rerun()


@st.dialog("Confirmar baja de personal")
def mostrar_confirmacion_baja_personal(nombre):
    st.warning(f"¿Confirmas que deseas dar de baja a {nombre}?")
    col_confirmar, col_cancelar = st.columns(2)
    with col_confirmar:
        confirmar = st.button(
            "Sí, dar de baja",
            key="confirmar_baja_personal",
            type="primary",
            use_container_width=True,
        )
    with col_cancelar:
        cancelar = st.button(
            "Cancelar",
            key="cancelar_baja_personal",
            type="secondary",
            use_container_width=True,
        )

    if cancelar:
        st.session_state.pop("gestion_seleccion_activa", None)
        st.rerun()
    if not confirmar:
        return
    if nombre not in lista_lavadores:
        st.error("Este personal ya no está registrado.")
        return

    lista_actualizada = [
        nombre_personal
        for nombre_personal in lista_lavadores
        if nombre_personal != nombre
    ]
    guardar_catalogo(ARCHIVO_LAVADORES, lista_actualizada)
    for clave in ("input_l1", "input_l2", "input_l3"):
        if st.session_state.get(clave) == nombre:
            st.session_state[clave] = "-"
    st.session_state.pop("gestion_seleccion_activa", None)
    notificar_exito(
        f"Lavador {nombre} dado de baja correctamente.",
        "eliminacion",
        destino="personal",
    )
    st.rerun()


lista_pagos = ["Efectivo", "Tarjeta", "Yape", "Plin", "Transferencia"]
l1_opts = ["-", *lista_lavadores]
l2_opts = ["-", *lista_lavadores]
l3_opts = ["-", *lista_lavadores]

if "edit_id" not in st.session_state: st.session_state.edit_id = None
if "registro_cargado_id" not in st.session_state: st.session_state.registro_cargado_id = None
if "revision_tabla_fecha" not in st.session_state: st.session_state.revision_tabla_fecha = 0
if "input_tipo" not in st.session_state: st.session_state.input_tipo = "Auto"
if "input_placa" not in st.session_state: st.session_state.input_placa = ""
if "input_servicio" not in st.session_state: st.session_state.input_servicio = "Lavado Simple"
if "input_monto" not in st.session_state: st.session_state.input_monto = "16.00"
if "input_pago" not in st.session_state: st.session_state.input_pago = "Efectivo"
if "input_l1" not in st.session_state: st.session_state.input_l1 = "Carlos"
if "input_l2" not in st.session_state: st.session_state.input_l2 = "-"
if "input_l3" not in st.session_state: st.session_state.input_l3 = "-"
if "input_motivo" not in st.session_state: st.session_state.input_motivo = "-"
if "input_precio_gasto" not in st.session_state: st.session_state.input_precio_gasto = "-"
if "input_metodo_gasto" not in st.session_state: st.session_state.input_metodo_gasto = "-"
if "input_ruc_cliente" not in st.session_state: st.session_state.input_ruc_cliente = ""
if "input_razon_social_cliente" not in st.session_state: st.session_state.input_razon_social_cliente = ""
if "input_numero_cliente" not in st.session_state: st.session_state.input_numero_cliente = ""
if "input_correo_cliente" not in st.session_state: st.session_state.input_correo_cliente = ""
if (
    st.session_state.input_pago not in lista_pagos
    and not (
        st.session_state.edit_id
        and st.session_state.input_pago == "Yape/Plin"
    )
):
    st.session_state.input_pago = "Efectivo"

if st.session_state.input_tipo not in lista_tipos:
    st.session_state.input_tipo = lista_tipos[0] if lista_tipos else ""
if st.session_state.input_servicio not in lista_servicios:
    st.session_state.input_servicio = lista_servicios[0] if lista_servicios else ""
if st.session_state.input_l1 not in l1_opts:
    st.session_state.input_l1 = "Carlos" if "Carlos" in lista_lavadores else (lista_lavadores[0] if lista_lavadores else "-")
if st.session_state.input_l2 not in l2_opts:
    st.session_state.input_l2 = "-"
if st.session_state.input_l3 not in l3_opts:
    st.session_state.input_l3 = "-"

def restablecer_campos_formulario():
    st.session_state.input_tipo = lista_tipos[0] if lista_tipos else ""
    st.session_state.input_placa = ""
    st.session_state.input_servicio = lista_servicios[0] if lista_servicios else ""
    st.session_state.input_monto = "16.00"
    st.session_state.input_pago = "Efectivo"
    st.session_state.input_l1 = "Carlos" if "Carlos" in lista_lavadores else (lista_lavadores[0] if lista_lavadores else "-")
    st.session_state.input_l2 = "-"
    st.session_state.input_l3 = "-"
    st.session_state.input_motivo = "-"
    st.session_state.input_precio_gasto = "-"
    st.session_state.input_metodo_gasto = "-"
    st.session_state.input_ruc_cliente = ""
    st.session_state.input_razon_social_cliente = ""
    st.session_state.input_numero_cliente = ""
    st.session_state.input_correo_cliente = ""


if st.session_state.pop("resetear_formulario", False):
    restablecer_campos_formulario()


def limpiar_formulario():
    st.session_state.edit_id = None
    st.session_state.resetear_formulario = True

def cambiar_fecha_tabla():
    st.session_state.revision_tabla_fecha += 1
    if st.session_state.edit_id:
        st.session_state.edit_id = None
        st.session_state.registro_cargado_id = None
        st.session_state.resetear_formulario = True

def cargar_fila_en_formulario(fila_data):
    registro_id = str(fila_data['ID'])
    if st.session_state.registro_cargado_id == registro_id:
        return

    st.session_state.registro_cargado_id = registro_id
    st.session_state.edit_id = registro_id
    st.session_state.input_tipo = fila_data['Tipo'] if fila_data['Tipo'] in lista_tipos else lista_tipos[0]
    st.session_state.input_placa = str(fila_data['Placa'])
    st.session_state.input_servicio = fila_data['Servicio'] if fila_data['Servicio'] in lista_servicios else lista_servicios[0]
    st.session_state.input_monto = str(fila_data['Monto']).replace("S/", "").strip()
    metodo_pago = str(fila_data['Pago']).strip()
    st.session_state.input_pago = (
        metodo_pago
        if metodo_pago in {*lista_pagos, "Yape/Plin"}
        else "Efectivo"
    )
    st.session_state.input_l1 = fila_data['Lavador 1'] if fila_data['Lavador 1'] in l1_opts else (
        "Carlos" if "Carlos" in lista_lavadores else (lista_lavadores[0] if lista_lavadores else "-")
    )
    st.session_state.input_l2 = fila_data['Lavador 2'] if fila_data['Lavador 2'] in l2_opts else "-"
    st.session_state.input_l3 = fila_data['Lavador 3'] if fila_data['Lavador 3'] in l3_opts else "-"
    st.session_state.input_motivo = "" if pd.isna(fila_data['Motivo Gasto']) or str(fila_data['Motivo Gasto']).strip() == "-" else str(fila_data['Motivo Gasto'])
    precio_gasto = str(fila_data['Precio Gasto']).replace("S/", "").strip()
    st.session_state.input_precio_gasto = "" if precio_gasto.casefold() in {"", "-", "nan", "none"} else precio_gasto
    metodo_gasto = str(fila_data['Método Gasto']).strip()
    st.session_state.input_metodo_gasto = metodo_gasto if metodo_gasto in lista_pagos else "-"
    valor_ruc = str(fila_data['RUC/DNI'])
    st.session_state.input_ruc_cliente = "" if valor_ruc.casefold() in {"", "-", "nan", "none"} else valor_ruc
    razon_social = str(fila_data['Razón Social/Nombre'])
    st.session_state.input_razon_social_cliente = "" if razon_social.casefold() in {"", "-", "nan", "none"} else razon_social
    numero_cliente = str(fila_data['Número Cliente'])
    st.session_state.input_numero_cliente = "" if numero_cliente.casefold() in {"", "-", "nan", "none"} else numero_cliente
    correo_cliente = str(fila_data['Correo Cliente'])
    st.session_state.input_correo_cliente = "" if correo_cliente.casefold() in {"", "-", "nan", "none"} else correo_cliente

# --- NAVEGACIÓN Y VISTAS ---
opciones_navegacion = ["Registros", "Gestión", "Control", "Finanzas"]
claves_navegacion = {
    "Registros": "registros",
    "Gestión": "gestion",
    "Control": "control",
    "Finanzas": "finanzas",
}
if st.session_state.get("menu_activo") not in opciones_navegacion:
    menu_anterior = st.session_state.get(
        "menu_activo", st.session_state.get("seccion_navegacion")
    )
    menu_anterior = {
        "Personal": "Gestión",
        "Ajustes": "Control",
        "Caja": "Finanzas",
    }.get(
        menu_anterior, menu_anterior
    )
    st.session_state.menu_activo = (
        menu_anterior if menu_anterior in opciones_navegacion else "Registros"
    )


def actualizar_menu_activo(opcion):
    st.session_state.menu_activo = opcion


st.sidebar.markdown(
    f'<div class="sidebar-logo-container"><img src="data:image/png;base64,{logo_base64}" alt="The Bunker Car Wash"></div>',
    unsafe_allow_html=True,
)
st.sidebar.markdown(
    """
    <style>
    [data-testid="stSidebar"] .st-key-navegacion_registros [data-testid^="stBaseButton-"],
    [data-testid="stSidebar"] .st-key-navegacion_gestion [data-testid^="stBaseButton-"],
    [data-testid="stSidebar"] .st-key-navegacion_control [data-testid^="stBaseButton-"],
    [data-testid="stSidebar"] .st-key-navegacion_finanzas [data-testid^="stBaseButton-"] {
        min-height: 46px;
        border-radius: 12px;
        justify-content: flex-start;
        padding-left: 1rem;
        font-weight: 600;
        transition: background-color 140ms ease, border-color 140ms ease;
    }
    [data-testid="stSidebar"] .st-key-navegacion_registros [data-testid="stBaseButton-secondary"],
    [data-testid="stSidebar"] .st-key-navegacion_gestion [data-testid="stBaseButton-secondary"],
    [data-testid="stSidebar"] .st-key-navegacion_control [data-testid="stBaseButton-secondary"],
    [data-testid="stSidebar"] .st-key-navegacion_finanzas [data-testid="stBaseButton-secondary"] {
        border: 1px solid rgba(128, 128, 128, 0.35);
    }
    </style>
    """,
    unsafe_allow_html=True,
)

for opcion in opciones_navegacion:
    st.sidebar.button(
        opcion,
        key=f"navegacion_{claves_navegacion[opcion]}",
        type="primary" if st.session_state.menu_activo == opcion else "secondary",
        use_container_width=True,
        on_click=actualizar_menu_activo,
        args=(opcion,),
    )

seccion = st.session_state.menu_activo
if "admin_autorizado" not in st.session_state:
    st.session_state.admin_autorizado = False

is_admin = st.session_state.admin_autorizado
is_supervisor = str(st.session_state.get("rol_usuario", "")).startswith(
    "Supervisor"
)
df_registros = consultar_dataframe()
USUARIOS_SUPERVISORES = sorted(
    nombre
    for nombre, cuenta in usuarios.items()
    if cuenta["activo"]
    and cuenta["rol"] == "Supervisor"
)

def mostrar_menu_usuario():
    rol_perfil = (
        "Administrador"
        if st.session_state.rol_usuario == "Administrador"
        else "Supervisor"
        if st.session_state.rol_usuario.startswith("Supervisor")
        else "Lavador"
    )
    with st.popover(
        ":material/account_circle:",
        key="menu_usuario",
        help="Cuenta de usuario",
        width=240,
    ):
        st.markdown(
            '<div class="profile-menu-avatar"><span aria-hidden="true">account_circle</span></div>',
            unsafe_allow_html=True,
        )
        st.markdown(
            f'<p class="profile-menu-label">Usuario</p>'
            f'<p class="profile-menu-user">'
            f'{escape(st.session_state.usuario_actual)}</p>'
            f'<p class="profile-menu-label">Rol</p>'
            f'<p class="profile-menu-role">{escape(rol_perfil)}</p>',
            unsafe_allow_html=True,
        )
        st.divider()
        if st.button(
            "Cerrar sesión",
            key="cerrar_sesion",
            type="secondary",
            use_container_width=True,
        ):
            st.session_state.login_completado = False
            st.session_state.admin_autorizado = False
            st.session_state.menu_activo = "Registros"
            st.session_state.pop("_registros_sesion", None)
            st.session_state.pop("_alcance_registros_sesion", None)
            st.session_state.pop("_usuarios_sesion", None)
            st.session_state.pop("_usuarios_eliminados_sesion", None)
            for key in (
                "usuario_actual",
                "rol_usuario",
                "usuario_login",
                "contrasena_login",
                "password_admin",
            ):
                st.session_state.pop(key, None)
            st.rerun()

mostrar_menu_usuario()

if seccion != "Registros":
    _, logo_columna = st.columns([8, 2])
    with logo_columna:
        st.markdown(
            f'<div class="main-header-logo"><img src="data:image/png;base64,{logo_base64}" alt="The Bunker Car Wash"></div>',
            unsafe_allow_html=True,
        )

if (
    seccion != "Registros"
    and not is_admin
    and not (seccion in {"Gestión", "Control"} and is_supervisor)
):
    st.title(construir_titulo_seccion(seccion))
    st.error("Acceso restringido. Se requiere una cuenta de Administrador.")
    st.stop()

@st.dialog("Confirmar eliminación definitiva")
def mostrar_confirmacion_eliminar_supervisor(nombre):
    st.warning(
        "¿Desea eliminar definitivamente los registros y el historial de este supervisor?"
    )
    clave = hashlib.sha256(nombre.casefold().encode("utf-8")).hexdigest()[:12]
    col_confirmar, col_cancelar = st.columns(2)
    with col_confirmar:
        confirmar = st.button(
            "Eliminar definitivamente",
            key=f"confirmar_eliminar_supervisor_{clave}",
            type="primary",
            use_container_width=True,
        )
    with col_cancelar:
        cancelar = st.button(
            "Cancelar",
            key=f"cancelar_eliminar_supervisor_{clave}",
            type="secondary",
            use_container_width=True,
        )

    if cancelar:
        st.rerun()
    if not confirmar:
        return

    cuenta = usuarios.get(nombre)
    if (
        cuenta is None
        or cuenta["activo"]
        or cuenta["rol"] != "Supervisor"
    ):
        st.error("El supervisor ya no está disponible para eliminación.")
        return

    usuarios_actualizados = {
        nombre_usuario: detalle.copy()
        for nombre_usuario, detalle in usuarios.items()
    }
    del usuarios_actualizados[nombre]
    eliminados_actualizados = set(usuarios_eliminados_definitivamente)
    eliminados_actualizados.discard(nombre.casefold())

    try:
        guardar_eliminacion_supervisor(
            usuarios_actualizados,
            nombre,
            eliminados_actualizados,
        )
    except (OSError, psycopg2.Error) as error:
        st.error(f"No se pudo completar la eliminación definitiva: {error}")
    else:
        st.session_state.pop("supervisor_detalle_finanzas", None)
        st.session_state.pop("responsable_ajustes", None)
        st.session_state["mensaje_gestion_usuarios"] = (
            f"Se eliminaron la cuenta y los registros de {nombre}."
        )
        st.rerun()

if seccion == "Finanzas":
    st.title(construir_titulo_seccion(seccion))
    ahora_finanzas = datetime.now(ZONA_HORARIA_PERU)
    nombres_meses_finanzas = [
        "Enero",
        "Febrero",
        "Marzo",
        "Abril",
        "Mayo",
        "Junio",
        "Julio",
        "Agosto",
        "Septiembre",
        "Octubre",
        "Noviembre",
        "Diciembre",
    ]
    st.subheader("Balance de Ingresos")
    fecha_referencia_balance = seleccionar_fecha(
        "Fecha de referencia",
        "fecha_referencia_balance_finanzas",
        ahora_finanzas.date(),
    )
    fecha_referencia_mes = fecha_referencia_balance
    mes_referencia_finanzas = nombres_meses_finanzas[
        fecha_referencia_mes.month - 1
    ]
    periodo_mensual_finanzas = (
        f"{mes_referencia_finanzas} {fecha_referencia_mes.year}"
    )
    registros_balance_mes = filtrar_registros_por_mes(
        df_registros, fecha_referencia_mes
    )
    ingresos_mes, _ = desglosar_metodos_financieros(
        registros_balance_mes
    )
    ganancia_mes = sum(ingresos_mes.values())

    registros_balance_ingresos_dia = filtrar_registros_por_fecha(
        df_registros, fecha_referencia_balance
    )
    ingresos_dia, _ = desglosar_metodos_financieros(
        registros_balance_ingresos_dia
    )
    ganancia_dia = sum(ingresos_dia.values())

    def renderizar_desglose_metodos(
        titulo, fecha, montos, periodo_tipo, categoria
    ):
        orden_metodos = [
            metodo
            for metodo in (
                "Efectivo",
                "Yape",
                "Plin",
                "Tarjeta",
                "Transferencia",
            )
            if metodo in montos
        ]
        orden_metodos.extend(
            sorted(
                (
                    metodo
                    for metodo, monto in montos.items()
                    if metodo not in orden_metodos and monto != 0
                ),
                key=str.casefold,
            )
        )
        clases_metodo = {
            "efectivo": "efectivo",
            "yape": "yape",
            "plin": "plin",
            "transferencia": "transferencia",
            "tarjeta": "tarjeta",
        }
        tarjetas = "".join(
            '<div class="reportes-kpi-card '
            f'{"gasto" if categoria == "gasto" else "supervisor"} '
            f'finanzas-metodo-card finanzas-metodo-{clases_metodo.get(metodo.casefold(), "otros")}">'
            '<div class="reportes-kpi-acento"></div>'
            '<div class="reportes-kpi-icono" aria-hidden="true">'
            f'{"payments" if categoria == "ingreso" else "account_balance_wallet"}'
            '</div>'
            f'<div class="reportes-kpi-etiqueta">{escape(metodo)}</div>'
            f'<div class="reportes-kpi-valor">S/ {montos[metodo]:.2f}</div>'
            '<div class="reportes-kpi-unidad">'
            f'{"ingresos" if categoria == "ingreso" else "gastos"}'
            '</div></div>'
            for metodo in orden_metodos
        )
        periodo = (
            fecha.strftime("%d/%m/%Y")
            if periodo_tipo == "dia"
            else f"{nombres_meses_finanzas[fecha.month - 1]} {fecha.year}"
        )
        st.markdown(
            '<section class="reportes-kpi-panel">'
            '<div class="reportes-kpi-header">'
            f'<span class="reportes-kpi-heading">{escape(titulo)}</span>'
            f'<span class="reportes-kpi-periodo">{periodo}</span>'
            '</div>'
            f'<div class="finanzas-kpi-grid">{tarjetas}</div>'
            '</section>',
            unsafe_allow_html=True,
        )

    st.markdown(
        '<div class="finanzas-balance-grid">'
        '<div class="finanzas-balance-card">'
        '<div class="finanzas-balance-etiqueta">Ingresos netos del día</div>'
        f'<div class="finanzas-balance-valor">S/ {ganancia_dia:.2f}</div>'
        '</div>'
        '<div class="finanzas-balance-card">'
        '<div class="finanzas-balance-etiqueta">Ingresos netos del mes</div>'
        f'<div class="finanzas-balance-valor">S/ {ganancia_mes:.2f}</div>'
        '</div>'
        '</div>',
        unsafe_allow_html=True,
    )
    with st.expander("Ver ingresos por método (Día)", expanded=False):
        renderizar_desglose_metodos(
            "Ingresos por método · Día seleccionado",
            fecha_referencia_balance,
            ingresos_dia,
            "dia",
            "ingreso",
        )
    with st.expander("Ver ingresos por método (Mes)", expanded=False):
        renderizar_desglose_metodos(
            f"Ingresos por método · {periodo_mensual_finanzas}",
            fecha_referencia_mes,
            ingresos_mes,
            "mes",
            "ingreso",
        )

    st.subheader("Balance de Gastos")
    fecha_referencia_gastos = seleccionar_fecha(
        "Fecha de referencia",
        "fecha_referencia_gastos_finanzas",
        ahora_finanzas.date(),
    )
    fecha_referencia_mes_gastos = fecha_referencia_gastos
    periodo_mensual_gastos = (
        f"{nombres_meses_finanzas[fecha_referencia_mes_gastos.month - 1]} "
        f"{fecha_referencia_mes_gastos.year}"
    )
    registros_balance_gastos_mes = filtrar_registros_por_mes(
        df_registros, fecha_referencia_mes_gastos
    )
    _, gastos_mes = desglosar_metodos_financieros(
        registros_balance_gastos_mes
    )
    gasto_mes = sum(gastos_mes.values())
    registros_balance_gastos_dia = filtrar_registros_por_fecha(
        df_registros, fecha_referencia_gastos
    )
    _, gastos_dia = desglosar_metodos_financieros(
        registros_balance_gastos_dia
    )
    gasto_dia = sum(gastos_dia.values())
    st.markdown(
        '<div class="finanzas-balance-grid">'
        '<div class="finanzas-balance-card gasto">'
        '<div class="finanzas-balance-etiqueta">Gastos netos del día</div>'
        f'<div class="finanzas-balance-valor">S/ {gasto_dia:.2f}</div>'
        '</div>'
        '<div class="finanzas-balance-card gasto">'
        '<div class="finanzas-balance-etiqueta">Gastos netos del mes</div>'
        f'<div class="finanzas-balance-valor">S/ {gasto_mes:.2f}</div>'
        '</div>'
        '</div>',
        unsafe_allow_html=True,
    )
    with st.expander("Ver método de gasto (Día)", expanded=False):
        renderizar_desglose_metodos(
            "Gastos por método · Día seleccionado",
            fecha_referencia_gastos,
            gastos_dia,
            "dia",
            "gasto",
        )
    with st.expander("Ver método de gasto (Mes)", expanded=False):
        renderizar_desglose_metodos(
            f"Gastos por método · {periodo_mensual_gastos}",
            fecha_referencia_mes_gastos,
            gastos_mes,
            "mes",
            "gasto",
        )

    st.divider()
    with st.expander("Rendimiento de ingresos por supervisor", expanded=False):
        col_periodo_finanzas, col_fecha_finanzas, _ = st.columns(
            [1.2, 1.5, 3.3]
        )
        with col_periodo_finanzas:
            periodo_finanzas = st.selectbox(
                "Periodo del desglose",
                ["Día", "Mes"],
                key="periodo_desglose_finanzas",
            )
        with col_fecha_finanzas:
            fecha_finanzas = seleccionar_fecha(
                "Fecha de referencia",
                "fecha_desglose_finanzas",
                ahora_finanzas.date(),
            )

        ingresos_responsables = calcular_ingresos_por_responsable(
            df_registros, fecha_finanzas, periodo_finanzas
        )
        etiqueta_periodo_finanzas = (
            fecha_finanzas.strftime("%d/%m/%Y")
            if periodo_finanzas == "Día"
            else fecha_finanzas.strftime("%m/%Y")
        )
        tarjetas_finanzas = []
        for responsable, monto in ingresos_responsables.items():
            tarjetas_finanzas.append(
                '<div class="reportes-kpi-card supervisor">'
                '<div class="reportes-kpi-acento"></div>'
                '<div class="reportes-kpi-icono" aria-hidden="true">support_agent</div>'
                f'<div class="reportes-kpi-etiqueta">{escape(responsable)}</div>'
                f'<div class="reportes-kpi-valor">S/ {monto:.2f}</div>'
                '<div class="reportes-kpi-unidad">ingresos registrados</div>'
                '</div>'
            )
        tarjetas_finanzas.append(
            '<div class="reportes-kpi-card total">'
            '<div class="reportes-kpi-acento"></div>'
            '<div class="reportes-kpi-icono" aria-hidden="true">payments</div>'
            '<div class="reportes-kpi-etiqueta">Total del equipo</div>'
            f'<div class="reportes-kpi-valor">S/ {ingresos_responsables.sum():.2f}</div>'
            '<div class="reportes-kpi-unidad">ingresos del periodo</div>'
            '</div>'
        )
        st.markdown(
            '<section class="reportes-kpi-panel">'
            '<div class="reportes-kpi-header">'
            '<span class="reportes-kpi-heading">Ingresos por supervisor</span>'
            f'<span class="reportes-kpi-periodo">{periodo_finanzas} · '
            f'{etiqueta_periodo_finanzas}</span>'
            '</div>'
            f'<div class="finanzas-kpi-grid">{"".join(tarjetas_finanzas)}</div>'
            '</section>',
            unsafe_allow_html=True,
        )

    with st.expander("Detalle de ingresos del supervisor", expanded=False):
        col_usuario_detalle, col_periodo_detalle, col_fecha_detalle, _ = st.columns(
            [1.4, 1.0, 1.5, 2.1]
        )
        with col_usuario_detalle:
            supervisor_detalle = st.selectbox(
                "Selecciona un usuario",
                USUARIOS_SUPERVISORES,
                key="supervisor_detalle_finanzas",
                disabled=not USUARIOS_SUPERVISORES,
            )
        with col_periodo_detalle:
            periodo_detalle_finanzas = st.selectbox(
                "Periodo del detalle",
                ["Día", "Mes"],
                key="periodo_detalle_supervisor_finanzas",
            )
        with col_fecha_detalle:
            fecha_detalle_finanzas = seleccionar_fecha(
                "Fecha del detalle",
                "fecha_detalle_finanzas",
                ahora_finanzas.date(),
            )

        registros_supervisor_detalle = consultar_dataframe(
            supervisor=supervisor_detalle,
        )
        if periodo_detalle_finanzas == "Día":
            registros_supervisor_periodo = filtrar_registros_por_fecha(
                registros_supervisor_detalle, fecha_detalle_finanzas
            )
            tipo_periodo_detalle = "dia"
            etiqueta_periodo_detalle = fecha_detalle_finanzas.strftime("%d/%m/%Y")
            sufijo_periodo_detalle = "del día"
        else:
            registros_supervisor_periodo = filtrar_registros_por_mes(
                registros_supervisor_detalle, fecha_detalle_finanzas
            )
            tipo_periodo_detalle = "mes"
            etiqueta_periodo_detalle = (
                f"{nombres_meses_finanzas[fecha_detalle_finanzas.month - 1]} "
                f"{fecha_detalle_finanzas.year}"
            )
            sufijo_periodo_detalle = "del mes"

        ingresos_detalle_metodos, gastos_detalle_metodos = (
            desglosar_metodos_financieros(registros_supervisor_periodo)
        )
        ingreso_periodo_detalle = sum(ingresos_detalle_metodos.values())
        gasto_periodo_detalle = sum(gastos_detalle_metodos.values())
        tarjetas_detalle_supervisor = (
            '<div class="reportes-kpi-card supervisor">'
            '<div class="reportes-kpi-acento"></div>'
            '<div class="reportes-kpi-icono" aria-hidden="true">payments</div>'
            f'<div class="reportes-kpi-etiqueta">Ingresos netos '
            f'{sufijo_periodo_detalle}</div>'
            f'<div class="reportes-kpi-valor">S/ {ingreso_periodo_detalle:.2f}</div>'
            f'<div class="reportes-kpi-unidad">{etiqueta_periodo_detalle}</div>'
            '</div>'
            '<div class="reportes-kpi-card gasto">'
            '<div class="reportes-kpi-acento"></div>'
            '<div class="reportes-kpi-icono" aria-hidden="true">'
            'account_balance_wallet</div>'
            f'<div class="reportes-kpi-etiqueta">Gastos netos '
            f'{sufijo_periodo_detalle}</div>'
            f'<div class="reportes-kpi-valor">S/ {gasto_periodo_detalle:.2f}</div>'
            f'<div class="reportes-kpi-unidad">{etiqueta_periodo_detalle}</div>'
            '</div>'
        )
        st.markdown(
            '<section class="reportes-kpi-panel">'
            '<div class="reportes-kpi-header">'
            f'<span class="reportes-kpi-heading">{escape(supervisor_detalle or "")}</span>'
            f'<span class="reportes-kpi-periodo">{periodo_detalle_finanzas} · '
            f'{etiqueta_periodo_detalle}</span>'
            '</div>'
            f'<div class="finanzas-kpi-grid">{tarjetas_detalle_supervisor}</div>'
            '</section>',
            unsafe_allow_html=True,
        )
        renderizar_desglose_metodos(
            f"Ingresos por método · {supervisor_detalle} · "
            f"{periodo_detalle_finanzas.lower()}",
            fecha_detalle_finanzas,
            ingresos_detalle_metodos,
            tipo_periodo_detalle,
            "ingreso",
        )
        renderizar_desglose_metodos(
            f"Gastos por método · {supervisor_detalle} · "
            f"{periodo_detalle_finanzas.lower()}",
            fecha_detalle_finanzas,
            gastos_detalle_metodos,
            tipo_periodo_detalle,
            "gasto",
        )

    mostrar_notificacion_exito()
    st.stop()

if seccion == "Gestión":
    st.title(construir_titulo_seccion(seccion))
    st.caption("Administra el personal, los tipos de vehículos y los servicios disponibles.")
    st.markdown(
        """
        <style>
        .st-key-tabla_gestion_personal,
        .st-key-tabla_gestion_vehiculo,
        .st-key-tabla_gestion_servicio {
            gap: 0 !important;
        }
        .st-key-encabezado_tabla_gestion_personal,
        .st-key-encabezado_tabla_gestion_vehiculo,
        .st-key-encabezado_tabla_gestion_servicio {
            margin-bottom: 1.5rem !important;
        }
        .gestion-catalog-table-header {
            padding: 0.65rem 0.9rem;
            background: rgba(128, 139, 150, 0.12);
            border-bottom: 1px solid rgba(128, 139, 150, 0.24);
            color: #c5ccd4;
            font-size: 0.9rem;
            font-weight: 600;
        }
        [class*="st-key-fila_gestion_"] button {
            justify-content: flex-start !important;
            text-align: left !important;
            border-radius: 0 !important;
            border-top: 0 !important;
            border-color: rgba(128, 139, 150, 0.2) !important;
            background: rgba(20, 24, 32, 0.72) !important;
            color: #e5e7eb !important;
            -webkit-text-fill-color: #e5e7eb !important;
        }
        [class*="st-key-fila_gestion_"] button > div {
            justify-content: flex-start !important;
            width: 100%;
        }
        [class*="st-key-fila_gestion_"] button span {
            display: block;
            width: 100%;
            text-align: left !important;
        }
        [class*="st-key-fila_gestion_"] button:hover {
            background: rgba(179, 38, 45, 0.16) !important;
            border-color: rgba(179, 38, 45, 0.55) !important;
        }
        [class*="st-key-fila_gestion_"] button[kind="primary"] {
            background: #b3262d !important;
            border-color: #b3262d !important;
            color: #ffffff !important;
            -webkit-text-fill-color: #ffffff !important;
            box-shadow: none !important;
        }
        [class*="st-key-fila_gestion_"] button[kind="primary"]:hover {
            background: #b3262d !important;
            border-color: #b3262d !important;
        }
        </style>
        """,
        unsafe_allow_html=True,
    )
    catalogos_gestion = (
        (
            "personal",
            "Lavadores",
            "Control de Lavadores",
            lista_lavadores,
            "personas registradas",
            "Agregar Personal +",
            "abrir_agregar_personal",
            "Eliminar Persona",
            "dar_baja_personal_button",
        ),
        (
            "vehiculo",
            "Vehículos",
            "Tipos de Vehículos",
            lista_tipos,
            "tipos registrados",
            "Agregar Tipo +",
            "abrir_agregar_tipo_vehiculo",
            "Eliminar Vehículo",
            "quitar_tipo_button",
        ),
        (
            "servicio",
            "Servicio",
            "Tipos de Servicios",
            lista_servicios,
            "servicios registrados",
            "Agregar Servicio +",
            "abrir_agregar_servicio",
            "Eliminar Servicio",
            "quitar_servicio_button",
        ),
    )
    elementos_por_tipo = {
        "personal": lista_lavadores,
        "vehiculo": lista_tipos,
        "servicio": lista_servicios,
    }
    for (
        tipo_catalogo,
        titulo,
        titulo_tabla,
        elementos,
        etiqueta,
        etiqueta_agregar,
        clave_agregar,
        etiqueta_quitar,
        clave_quitar,
    ) in catalogos_gestion:
        with st.container(key=f"gestion_{tipo_catalogo}_panel", border=True):
            mostrar_notificacion_exito(tipo_catalogo)
            col_titulo, col_agregar, col_quitar = st.columns([3, 1.25, 1.6])
            with col_titulo:
                st.subheader(titulo)
                st.caption(f"{len(elementos)} {etiqueta}")
            with col_agregar:
                if st.button(
                    etiqueta_agregar,
                    key=clave_agregar,
                    type="primary",
                    use_container_width=True,
                ):
                    mostrar_formulario_catalogo(tipo_catalogo)

            seleccion_activa = st.session_state.get("gestion_seleccion_activa")
            if (
                not isinstance(seleccion_activa, dict)
                and seleccion_activa is not None
            ):
                st.session_state.pop("gestion_seleccion_activa", None)
                seleccion_activa = None
            if (
                seleccion_activa is not None
                and (
                    seleccion_activa.get("tipo") not in elementos_por_tipo
                    or seleccion_activa.get("elemento")
                    not in elementos_por_tipo.get(
                        seleccion_activa.get("tipo"), []
                    )
                )
            ):
                st.session_state.pop("gestion_seleccion_activa", None)
                seleccion_activa = None

            with col_quitar:
                if st.button(
                    etiqueta_quitar,
                    key=clave_quitar,
                    type="secondary",
                    use_container_width=True,
                ):
                    seleccion_activa = st.session_state.get(
                        "gestion_seleccion_activa"
                    )
                    if (
                        seleccion_activa is None
                        or seleccion_activa.get("tipo") != tipo_catalogo
                    ):
                        mensajes_sin_seleccion = {
                            "personal": "Selecciona un personal antes de darlo de baja.",
                            "vehiculo": "Se necesita seleccionar un tipo de vehículo.",
                            "servicio": "Se necesita seleccionar un servicio.",
                        }
                        st.error(mensajes_sin_seleccion[tipo_catalogo])
                    elif tipo_catalogo == "personal":
                        mostrar_confirmacion_baja_personal(
                            seleccion_activa["elemento"]
                        )
                    elif len(elementos) <= 1:
                        st.error(
                            "No se puede quitar el último tipo de vehículo."
                            if tipo_catalogo == "vehiculo"
                            else "No se puede quitar el último servicio."
                        )
                    else:
                        elemento_seleccionado = seleccion_activa["elemento"]
                        elementos_actualizados = [
                            elemento
                            for elemento in elementos
                            if elemento != elemento_seleccionado
                        ]
                        archivo_catalogo = (
                            ARCHIVO_VEHICULOS
                            if tipo_catalogo == "vehiculo"
                            else ARCHIVO_SERVICIOS
                        )
                        guardar_catalogo(archivo_catalogo, elementos_actualizados)
                        clave_input = (
                            "input_tipo"
                            if tipo_catalogo == "vehiculo"
                            else "input_servicio"
                        )
                        if st.session_state.get(clave_input) == elemento_seleccionado:
                            st.session_state[clave_input] = elementos_actualizados[0]
                        st.session_state.pop("gestion_seleccion_activa", None)
                        notificar_exito(
                            f"{'Tipo de vehículo' if tipo_catalogo == 'vehiculo' else 'Servicio'} "
                            f"{elemento_seleccionado} quitado correctamente.",
                            "eliminacion",
                            destino=tipo_catalogo,
                        )
                        st.rerun()

            with st.container(
                key=f"tabla_gestion_{tipo_catalogo}", border=True
            ):
                with st.container(
                    key=f"encabezado_tabla_gestion_{tipo_catalogo}"
                ):
                    st.markdown(
                        f'<div class="gestion-catalog-table-header">'
                        f"{escape(titulo_tabla)}</div>",
                        unsafe_allow_html=True,
                    )
                if elementos:
                    for elemento in elementos:
                        clave_elemento = hashlib.sha256(
                            elemento.casefold().encode("utf-8")
                        ).hexdigest()[:12]
                        clave_fila = (
                            f"fila_gestion_{tipo_catalogo}_{clave_elemento}"
                        )
                        seleccion_actual = st.session_state.get(
                            "gestion_seleccion_activa"
                        )
                        seleccionado = (
                            seleccion_actual is not None
                            and seleccion_actual.get("tipo") == tipo_catalogo
                            and seleccion_actual.get("elemento") == elemento
                        )
                        if st.button(
                            elemento,
                            key=clave_fila,
                            type="primary" if seleccionado else "secondary",
                            use_container_width=True,
                        ):
                            if seleccionado:
                                st.session_state.pop(
                                    "gestion_seleccion_activa", None
                                )
                            else:
                                st.session_state["gestion_seleccion_activa"] = {
                                    "tipo": tipo_catalogo,
                                    "elemento": elemento,
                                }
                            st.rerun()
                else:
                    st.info("No hay elementos registrados.")

        if tipo_catalogo == "vehiculo" and len(elementos) <= 1:
            st.caption("Debe quedar al menos un tipo de vehículo disponible.")
        elif tipo_catalogo == "servicio" and len(elementos) <= 1:
            st.caption("Debe quedar al menos un servicio disponible.")

    st.stop()

if seccion == "Control":
    st.title(construir_titulo_seccion(seccion))
    st.markdown("---")
    st.subheader(
        "Registro de Servicios y Gastos - Supervisor"
        if is_supervisor
        else "Registros de Servicios y Gastos"
    )
    if is_supervisor:
        col_fecha_ajustes, col_periodo_ajustes, _ = st.columns(
            [1.4, 1.1, 2.5]
        )
    else:
        col_fecha_ajustes, col_periodo_ajustes, col_responsable_ajustes = st.columns(
            [1, 1, 1.6]
        )
    with col_fecha_ajustes:
        fecha_registros_ajustes = seleccionar_fecha(
            "Fecha" if is_supervisor else "Ver Fecha",
            "fecha_registros_ajustes",
            datetime.now(ZONA_HORARIA_PERU).date(),
        )
    with col_periodo_ajustes:
        periodo_ajustes = st.selectbox(
            "Periodo", ["Día", "Mes"], key="periodo_ajustes"
        )
    if not is_supervisor:
        with col_responsable_ajustes:
            responsables_filtro = obtener_responsables_reporte(df_registros)
            responsable_ajustes = st.selectbox(
                "Responsable",
                ["Todos", *responsables_filtro],
                key="responsable_ajustes",
            )

    if periodo_ajustes == "Día":
        registros_periodo_ajustes = filtrar_registros_por_fecha(
            df_registros, fecha_registros_ajustes
        )
        registros_supervisor_ajustes = (
            consultar_dataframe(
                fecha=fecha_registros_ajustes,
                supervisor=st.session_state.usuario_actual,
            )
            if is_supervisor
            else None
        )
    else:
        registros_periodo_ajustes = filtrar_registros_por_mes(
            df_registros, fecha_registros_ajustes
        )
        registros_supervisor_ajustes = (
            consultar_dataframe(
                mes=fecha_registros_ajustes,
                supervisor=st.session_state.usuario_actual,
            )
            if is_supervisor
            else None
        )

    if is_supervisor:
        registros_ajustes = registros_supervisor_ajustes
    else:
        conteos_responsables = (
            registros_periodo_ajustes["Registrado por"]
            .map(normalizar_responsable)
            .value_counts()
        )
        registros_ajustes = registros_periodo_ajustes.copy()
        if responsable_ajustes != "Todos":
            registros_ajustes = registros_ajustes.loc[
                registros_ajustes["Registrado por"].map(normalizar_responsable)
                == responsable_ajustes
            ].copy()

    etiqueta_periodo_ajustes = (
        fecha_registros_ajustes.strftime("%d/%m/%Y")
        if periodo_ajustes == "Día"
        else fecha_registros_ajustes.strftime("%m/%Y")
    )
    if not is_supervisor:
        tarjetas_responsables = []
        for responsable in obtener_responsables_reporte(registros_periodo_ajustes):
            if responsable == "Administrador":
                icono, clase = "admin_panel_settings", "administrador"
            else:
                icono, clase = "support_agent", "supervisor"
            conteo = int(conteos_responsables.get(responsable, 0))
            tarjetas_responsables.append(
                f'<div class="reportes-kpi-card {clase}">'
                f'<div class="reportes-kpi-acento"></div>'
                f'<div class="reportes-kpi-icono" aria-hidden="true">{icono}</div>'
                f'<div class="reportes-kpi-etiqueta">{escape(responsable)}</div>'
                f'<div class="reportes-kpi-valor">{conteo}</div>'
                '<div class="reportes-kpi-unidad">reportes</div>'
                '</div>'
            )
        tarjetas_responsables.append(
            '<div class="reportes-kpi-card total">'
            '<div class="reportes-kpi-acento"></div>'
            '<div class="reportes-kpi-icono" aria-hidden="true">summarize</div>'
            '<div class="reportes-kpi-etiqueta">Total registrado</div>'
            f'<div class="reportes-kpi-valor">{len(registros_periodo_ajustes)}</div>'
            '<div class="reportes-kpi-unidad">reportes</div>'
            '</div>'
        )
    if registros_ajustes.empty:
        st.info("No hay registros que coincidan con los filtros seleccionados.")
    else:
        registros_ajustes = registros_ajustes.copy()
        registros_ajustes["Registrado por"] = registros_ajustes[
            "Registrado por"
        ].replace("", "No disponible")
        st.dataframe(
            crear_tabla_estilizada(registros_ajustes, incluir_usuario=True),
            use_container_width=True,
            hide_index=True,
        )

    if is_supervisor:
        ingresos_supervisor, gastos_supervisor = desglosar_metodos_financieros(
            registros_ajustes
        )
        clases_metodo_supervisor = {
            "efectivo": "efectivo",
            "yape": "yape",
            "plin": "plin",
            "tarjeta": "tarjeta",
            "transferencia": "transferencia",
        }
        periodo_resumen_supervisor = (
            f"{periodo_ajustes} · {etiqueta_periodo_ajustes}"
        )
        tarjetas_ingresos_supervisor = "".join(
            '<div class="reportes-kpi-card supervisor '
            'finanzas-metodo-card finanzas-metodo-'
            f'{clases_metodo_supervisor.get(metodo.casefold(), "otros")}">'
            '<div class="reportes-kpi-acento"></div>'
            '<div class="reportes-kpi-icono" aria-hidden="true">payments</div>'
            f'<div class="reportes-kpi-etiqueta">{escape(metodo)}</div>'
            f'<div class="reportes-kpi-valor">S/ {monto:.2f}</div>'
            '<div class="reportes-kpi-unidad">ingresos</div></div>'
            for metodo, monto in ingresos_supervisor.items()
        )
        tarjetas_gastos_supervisor = "".join(
            '<div class="reportes-kpi-card gasto '
            'finanzas-metodo-card finanzas-metodo-'
            f'{clases_metodo_supervisor.get(metodo.casefold(), "otros")}">'
            '<div class="reportes-kpi-acento"></div>'
            '<div class="reportes-kpi-icono" aria-hidden="true">'
            'account_balance_wallet</div>'
            f'<div class="reportes-kpi-etiqueta">{escape(metodo)}</div>'
            f'<div class="reportes-kpi-valor">S/ {monto:.2f}</div>'
            '<div class="reportes-kpi-unidad">gastos</div></div>'
            for metodo, monto in gastos_supervisor.items()
        )
        tarjetas_balance_supervisor = (
            '<div class="finanzas-balance-card">'
            '<div class="finanzas-balance-etiqueta">Ingresos totales del período</div>'
            f'<div class="finanzas-balance-valor">S/ '
            f'{sum(ingresos_supervisor.values()):.2f}</div></div>'
            '<div class="finanzas-balance-card gasto">'
            '<div class="finanzas-balance-etiqueta">Gastos totales del período</div>'
            f'<div class="finanzas-balance-valor">S/ '
            f'{sum(gastos_supervisor.values()):.2f}</div></div>'
        )
        st.markdown(
            '<section class="reportes-kpi-panel">'
            '<div class="reportes-kpi-header">'
            '<span class="reportes-kpi-heading">Balance del período</span>'
            f'<span class="reportes-kpi-periodo">{periodo_resumen_supervisor}</span>'
            '</div>'
            f'<div class="finanzas-balance-grid">{tarjetas_balance_supervisor}</div>'
            '</section>',
            unsafe_allow_html=True,
        )
        st.markdown(
            '<section class="reportes-kpi-panel">'
            '<div class="reportes-kpi-header">'
            '<span class="reportes-kpi-heading">Ingresos por método de pago</span>'
            f'<span class="reportes-kpi-periodo">{periodo_resumen_supervisor}</span>'
            '</div>'
            f'<div class="finanzas-kpi-grid">{tarjetas_ingresos_supervisor}</div>'
            '</section>',
            unsafe_allow_html=True,
        )
        st.markdown(
            '<section class="reportes-kpi-panel">'
            '<div class="reportes-kpi-header">'
            '<span class="reportes-kpi-heading">Gastos por método de pago</span>'
            f'<span class="reportes-kpi-periodo">{periodo_resumen_supervisor}</span>'
            '</div>'
            f'<div class="finanzas-kpi-grid">{tarjetas_gastos_supervisor}</div>'
            '</section>',
            unsafe_allow_html=True,
        )
        st.stop()

    ahora_ajustes = datetime.now(ZONA_HORARIA_PERU)
    registros_mes = filtrar_registros_del_mes(df_registros, ahora_ajustes)
    registros_dia = filtrar_registros_por_fecha(
        df_registros, fecha_registros_ajustes
    )
    if responsable_ajustes != "Todos":
        registros_mes = registros_mes.loc[
            registros_mes["Registrado por"].map(normalizar_responsable)
            == responsable_ajustes
        ].copy()
        registros_dia = registros_dia.loc[
            registros_dia["Registrado por"].map(normalizar_responsable)
            == responsable_ajustes
        ].copy()

    col_exportar_dia, col_exportar_mes = st.columns(2)
    with col_exportar_dia:
        st.download_button(
            "Exportar Excel por Día",
            data=generar_excel_registros(registros_dia),
            file_name=f"registros_{fecha_registros_ajustes:%Y-%m-%d}.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            key="exportar_excel_ajustes_dia",
            disabled=registros_dia.empty,
            use_container_width=True,
        )
    with col_exportar_mes:
        st.download_button(
            "Exportar Excel por Mes",
            data=generar_excel_registros(registros_mes),
            file_name=f"registros_{ahora_ajustes:%Y-%m}.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            key="exportar_excel_ajustes_mes",
            disabled=registros_mes.empty,
            use_container_width=True,
        )

    st.markdown(
        f'<section class="reportes-kpi-panel">'
        f'<div class="reportes-kpi-header">'
        f'<span class="reportes-kpi-heading">Reportes por responsable</span>'
        f'<span class="reportes-kpi-periodo">{periodo_ajustes} · '
        f'{etiqueta_periodo_ajustes}</span>'
        f'</div><div class="reportes-kpi-grid">'
        f'{"".join(tarjetas_responsables)}'
        '</div></section>',
        unsafe_allow_html=True,
    )

    st.markdown(
        """
        <style>
        .st-key-gestion_usuarios_panel {
            padding: 1.25rem;
            border: 1px solid rgba(173, 190, 174, 0.2);
            border-radius: 14px;
            background: rgba(13, 19, 15, 0.72);
            backdrop-filter: blur(14px);
            box-shadow: inset 0 1px 0 rgba(255, 255, 255, 0.04),
                0 12px 32px rgba(0, 0, 0, 0.22);
        }
        .st-key-gestion_usuarios_panel h3 {
            color: #f1f5ef;
        }
        </style>
        """,
        unsafe_allow_html=True,
    )
    with st.container(key="gestion_usuarios_panel"):
        st.subheader("Gestión y Control de Usuarios")
        st.caption("Consulta las cuentas registradas y administra los accesos al sistema.")

        mensaje_usuarios = st.session_state.pop("mensaje_gestion_usuarios", None)
        if mensaje_usuarios:
            st.success(mensaje_usuarios)
        usuarios_activos = [
            (nombre, cuenta)
            for nombre, cuenta in usuarios.items()
            if cuenta["activo"]
        ]
        usuarios_activos.sort(
            key=lambda elemento: (
                elemento[1]["rol"] != "Administrador",
                elemento[1]["rol"],
                elemento[0].casefold(),
            )
        )
        tabla_usuarios = pd.DataFrame(
            [
                {
                    "Nombre de usuario": nombre,
                    "Rol": cuenta["rol"],
                    "Último inicio de sesión": (
                        datetime.fromisoformat(cuenta["ultimo_inicio_sesion"]).strftime(
                            "%d/%m/%Y %H:%M"
                        )
                        if cuenta["ultimo_inicio_sesion"]
                        else "Sin registros"
                    ),
                }
                for nombre, cuenta in usuarios_activos
            ]
        )
        st.dataframe(
            tabla_usuarios,
            use_container_width=True,
            hide_index=True,
        )

        supervisores_inactivos = [
            nombre
            for nombre, cuenta in usuarios.items()
            if not cuenta["activo"] and cuenta["rol"] == "Supervisor"
        ]
        if supervisores_inactivos:
            st.markdown("### Supervisores dados de baja")
            st.caption(
                "Elimina definitivamente una cuenta y todos sus registros asociados."
            )
            st.markdown(
                """
                <style>
                [class*="st-key-supervisor_baja_"] [data-testid="stBaseButton-secondary"] {
                    width: 2.35rem;
                    min-width: 2.35rem;
                    height: 2.35rem;
                    padding: 0;
                    border: 0;
                    border-radius: 50%;
                    background: #c62828;
                    color: #ffffff;
                    font-size: 1.25rem;
                    font-weight: 700;
                }
                [class*="st-key-supervisor_baja_"] [data-testid="stBaseButton-secondary"]:hover {
                    border: 0;
                    background: #a61f1f;
                    color: #ffffff;
                }
                </style>
                """,
                unsafe_allow_html=True,
            )
            columnas_supervisores = st.columns(3)
            for indice, nombre in enumerate(supervisores_inactivos):
                clave = hashlib.sha256(
                    nombre.casefold().encode("utf-8")
                ).hexdigest()[:12]
                with columnas_supervisores[indice % len(columnas_supervisores)]:
                    with st.container(key=f"supervisor_baja_{clave}", border=True):
                        col_nombre, col_eliminar = st.columns([5, 1])
                        with col_nombre:
                            st.markdown(f"**{escape(nombre)}**")
                        with col_eliminar:
                            if st.button(
                                "×",
                                key=f"eliminar_supervisor_{clave}",
                                help="Eliminar definitivamente los registros del supervisor",
                                type="secondary",
                            ):
                                mostrar_confirmacion_eliminar_supervisor(nombre)
                        st.caption("Cuenta dada de baja")

        col_baja, col_editar, col_nuevo = st.columns([1.3, 1, 1])
        with col_baja:
            if st.button(
                "Dar de baja una cuenta",
                key="abrir_formulario_baja_usuario",
                type="secondary",
                use_container_width=True,
            ):
                mostrar_formulario_baja_usuario()

        with col_editar:
            if st.button(
                "Editar usuario",
                key="abrir_formulario_editar_usuario",
                type="secondary",
                use_container_width=True,
            ):
                mostrar_formulario_editar_usuario()

        with col_nuevo:
            if st.button(
                "Agregar Nuevo Usuario",
                key="abrir_formulario_nuevo_usuario",
                type="primary",
                use_container_width=True,
            ):
                mostrar_formulario_nuevo_usuario()

    st.stop()

with st.container(key="registros_hero"):
    st.markdown('<div id="registros-theme-marker"></div>', unsafe_allow_html=True)
    _, logo_columna = st.columns([8, 2])
    with logo_columna:
        st.markdown(
            f'<div class="main-header-logo"><img src="data:image/png;base64,{logo_base64}" alt="The Bunker Car Wash"></div>',
            unsafe_allow_html=True,
        )
    columnas_banner = st.columns([1, 3])
    col_fecha = columnas_banner[0]
    with col_fecha:
        fecha_ver = seleccionar_fecha(
            "Ver Fecha",
            "ver_fecha_registros",
            datetime.now(ZONA_HORARIA_PERU).date(),
            on_change=cambiar_fecha_tabla,
        )
st.title(construir_titulo_seccion("Registros"))
st.caption(f"Fecha: {fecha_ver.strftime('%d/%m/%Y')}")

# --- TABLA DE REGISTROS (PROCESADA PRIMERO PARA CARGAR LA SELECCIÓN) ---
df_registros_fecha = consultar_dataframe(
    fecha=fecha_ver,
)
df_filtrado = df_registros_fecha.copy()
df_exportacion = df_registros_fecha.copy()

if is_admin:
    st.subheader("Selecciona una fila para editar")

    with st.container(border=True):
        col_f1, col_f2 = st.columns(2)
        with col_f1:
            busqueda_placa = st.text_input("Buscar por Placa", key="busqueda_placa_input")
        with col_f2:
            filtro_pago = st.selectbox(
                "Filtrar por Método de Pago",
                ["Todos", *lista_pagos, "Yape/Plin"],
                key="filtro_pago_select",
            )
        
    if busqueda_placa:
        df_filtrado = df_filtrado[df_filtrado['Placa'].astype(str).str.contains(busqueda_placa, case=False, na=False)]
    if filtro_pago != "Todos":
        df_filtrado = df_filtrado[df_filtrado['Pago'] == filtro_pago]
    df_exportacion = df_filtrado.copy()

    clave_tabla = f"tabla_admin_{st.session_state.revision_tabla_fecha}"
    st.session_state.tabla_admin_key_activa = clave_tabla
    filas_seleccionadas = st.session_state.get(f"filas_{clave_tabla}", set())

    with st.container(border=True):
        evento_tabla = st.dataframe(
            crear_tabla_estilizada(
                df_filtrado,
                filas_seleccionadas,
            ),
            use_container_width=True,
            column_order=obtener_columnas_visibles_registros(df_filtrado),
            selection_mode="single-row",
            on_select=guardar_seleccion_tabla,
            key=clave_tabla,
            hide_index=True,
        )
    
    seleccion_tabla = (
        evento_tabla.get("selection", {})
        if evento_tabla is not None
        else {}
    )
    filas_seleccionadas = seleccion_tabla.get("rows", [])
    if filas_seleccionadas:
        indice_seleccionado = filas_seleccionadas[0]
        if len(df_filtrado) > indice_seleccionado:
            fila_data = df_filtrado.iloc[indice_seleccionado]
            cargar_fila_en_formulario(fila_data)
    elif st.session_state.edit_id:
        st.session_state.edit_id = None
        st.session_state.registro_cargado_id = None
        restablecer_campos_formulario()
    else:
        st.session_state.registro_cargado_id = None

    if df_filtrado.empty:
        if df_registros_fecha.empty:
            st.info("No hay registros para la fecha seleccionada.")
        else:
            st.info("No hay movimientos que coincidan con los filtros.")

if not is_admin and not df_registros_fecha.empty:
    with st.container(border=True):
        st.dataframe(
            crear_tabla_estilizada(
                df_registros_fecha,
            ),
            use_container_width=True,
            column_order=obtener_columnas_visibles_registros(
                df_registros_fecha
            ),
            hide_index=True,
            key="tabla_registros_supervisor",
        )
elif not is_admin:
    st.info(f"No hay registros para el {fecha_ver.strftime('%d/%m/%Y')}.")

st.download_button(
    "Exportar a Excel",
    data=generar_excel_registros(df_exportacion),
    file_name=(
        f"registros_{fecha_ver.strftime('%Y-%m-%d')}_bunker.xlsx"
    ),
    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    key="descargar_registros_excel",
    type="primary",
    use_container_width=True,
)

st.markdown("---")

# --- FORMULARIO DE REGISTRO / ACTUALIZACIÓN ---
titulo_form = "Actualizar Registro Seleccionado" if st.session_state.edit_id else "Registrar Nuevo Servicio / Gasto"
st.subheader(titulo_form)

if st.session_state.edit_id:
    st.warning(
        f"Modificando el registro de la placa: "
        f"**{st.session_state.input_placa}**. "
        "Modifica los campos y haz clic en Guardar Cambios."
    )

with st.container(border=True):
    st.markdown("### Vehículo y Pago")
    c1, c2, c3, c4, c5 = st.columns(5)
    opciones_pago_formulario = list(lista_pagos)
    if (
        st.session_state.edit_id
        and st.session_state.input_pago == "Yape/Plin"
    ):
        opciones_pago_formulario.append("Yape/Plin")
    
    with c1:
        st.selectbox("Tipo de Vehículo", lista_tipos, key="input_tipo")
    with c2:
        st.text_input("Placa / Identificador", placeholder="Ej. ABC-123", key="input_placa")
    with c3:
        st.selectbox(
            "Servicio",
            lista_servicios,
            key="input_servicio",
            accept_new_options=True,
        )
    with c4:
        st.text_input("Monto Cobrado (S/)", key="input_monto")
    with c5:
        st.selectbox(
            "Forma de Pago",
            opciones_pago_formulario,
            key="input_pago",
        )
        
with st.container(border=True):
    st.markdown("### Cliente")
    cliente_ruc, cliente_razon_social, cliente_numero, cliente_correo = st.columns(4)
    with cliente_ruc:
        st.text_input("RUC/DNI (opcional)", key="input_ruc_cliente")
    with cliente_razon_social:
        st.text_input(
            "Razón Social/Nombre (opcional)",
            key="input_razon_social_cliente",
        )
    with cliente_numero:
        st.text_input("Número Cliente (opcional)", key="input_numero_cliente")
    with cliente_correo:
        st.text_input("Correo Cliente (opcional)", key="input_correo_cliente")

with st.container(border=True):
    st.markdown("### Lavadores")
    l1, l2, l3 = st.columns(3)
    with l1:
        st.selectbox("Lavador 1", l1_opts, key="input_l1")
    with l2:
        st.selectbox("Lavador 2", l2_opts, key="input_l2")
    with l3:
        st.selectbox("Lavador 3", l3_opts, key="input_l3")
        
with st.container(border=True):
    st.markdown("### Gastos")
    g1, g2, g3 = st.columns(3)
    with g1:
        st.text_input("Motivo Gasto (opcional)", key="input_motivo")
    with g2:
        st.text_input("Precio Gasto (opcional)", key="input_precio_gasto")
    with g3:
        st.selectbox(
            "Método Gasto",
            ["-", *lista_pagos],
            key="input_metodo_gasto",
        )

col_acc1, col_acc2 = st.columns(2)
with col_acc1:
    texto_boton = "Guardar Cambios de Actualización" if st.session_state.edit_id else "Guardar Registro"
    btn_guardar = st.button(
        texto_boton, key="guardar_registro_btn", use_container_width=True, type="primary"
    )
with col_acc2:
    if st.session_state.edit_id:
        btn_eliminar = st.button(
            "Eliminar este Registro", key="eliminar_registro_btn", type="secondary",
            use_container_width=True
        )
    else:
        btn_eliminar = False

mostrar_notificacion_exito()

# Lógica de Guardar / Actualizar
if btn_guardar:
    campos_obligatorios = {
        "Tipo de Vehículo": st.session_state.input_tipo,
        "Placa / Identificador": st.session_state.input_placa,
        "Servicio": st.session_state.input_servicio,
        "Monto Cobrado (S/)": st.session_state.input_monto,
        "Forma de Pago": st.session_state.input_pago,
    }
    campos_vacios = [
        nombre
        for nombre, valor in campos_obligatorios.items()
        if valor is None or not str(valor).strip() or str(valor).strip() == "-"
    ]
    m_val = str(st.session_state.input_monto).strip()
    pg_val = str(st.session_state.input_precio_gasto).strip()
    errores_validacion = []

    if campos_vacios:
        errores_validacion.append(
            "Completa los campos obligatorios: " + ", ".join(campos_vacios) + "."
        )
    if m_val and not re.fullmatch(r"\d+(?:\.\d+)?", m_val):
        errores_validacion.append(
            "Monto Cobrado solo acepta números y un punto decimal."
        )
    if pg_val not in {"", "-"} and not re.fullmatch(r"\d+(?:\.\d+)?", pg_val):
        errores_validacion.append(
            "Precio Gasto solo acepta números y un punto decimal."
        )
    errores_validacion.extend(
        validar_campos_gasto(
            pg_val,
            st.session_state.input_metodo_gasto,
        )
    )
    lavador_1 = str(st.session_state.input_l1).strip()
    lavador_2 = str(st.session_state.input_l2).strip()
    lavador_3 = str(st.session_state.input_l3).strip()
    lavadores_seleccionados = [
        lavador
        for lavador in (lavador_1, lavador_2, lavador_3)
        if lavador != "-"
    ]
    if not lavadores_seleccionados:
        errores_validacion.append("Selecciona al menos un lavador para el servicio.")
    elif len({lavador.casefold() for lavador in lavadores_seleccionados}) < len(
        lavadores_seleccionados
    ):
        errores_validacion.append(
            "No puedes seleccionar el mismo lavador más de una vez."
        )

    if errores_validacion:
        st.error("No se pudo guardar:\n\n- " + "\n- ".join(errores_validacion))
    elif st.session_state.input_placa.strip() != "":
        monto_fmt = f"S/ {float(m_val):.2f}"
        gasto_fmt = (
            f"S/ {float(pg_val):.2f}"
            if pg_val not in {"", "-"}
            else "-"
        )
        
        if st.session_state.edit_id:
            cambios_registro = {
                "Tipo": st.session_state.input_tipo,
                "Placa": st.session_state.input_placa.upper(),
                "Servicio": st.session_state.input_servicio,
                "Pago": st.session_state.input_pago,
                "Monto": monto_fmt,
                "Lavador 1": lavador_1 or "-",
                "Lavador 2": lavador_2 or "-",
                "Lavador 3": lavador_3 or "-",
                "Motivo Gasto": st.session_state.input_motivo.strip() or "-",
                "Precio Gasto": gasto_fmt,
                "Método Gasto": st.session_state.input_metodo_gasto or "-",
                "RUC/DNI": st.session_state.input_ruc_cliente.strip() or "-",
                "Razón Social/Nombre": (
                    st.session_state.input_razon_social_cliente.strip() or "-"
                ),
                "Número Cliente": st.session_state.input_numero_cliente.strip() or "-",
                "Correo Cliente": st.session_state.input_correo_cliente.strip() or "-",
            }
            try:
                actualizado = actualizar_registro(
                    st.session_state.edit_id,
                    cambios_registro,
                )
            except (psycopg2.Error, ValueError) as error:
                st.error(f"No se pudo actualizar el registro: {error}")
            else:
                if actualizado:
                    refrescar_registros_sesion()
                    limpiar_formulario()
                    notificar_exito("¡Se actualizó correctamente!")
                    st.rerun()
                else:
                    st.error("No se encontró el registro a actualizar.")
        else:
            ahora_registro = datetime.now(ZONA_HORARIA_PERU)
            fecha_registro = (
                f"{fecha_ver:%Y/%m/%d} {ahora_registro:%H:%M}"
            )
            nuevo_registro = {
                "Tipo": st.session_state.input_tipo,
                "Placa": st.session_state.input_placa.upper(),
                "Servicio": st.session_state.input_servicio,
                "Pago": st.session_state.input_pago,
                "Monto": monto_fmt,
                "Fecha": fecha_registro,
                "Lavador 1": st.session_state.input_l1,
                "Lavador 2": st.session_state.input_l2,
                "Lavador 3": st.session_state.input_l3,
                "Motivo Gasto": st.session_state.input_motivo.strip() or "-",
                "Precio Gasto": gasto_fmt,
                "Método Gasto": st.session_state.input_metodo_gasto or "-",
                NOMBRE_COLUMNA_LOCAL: "Búnker 1",
                "RUC/DNI": (
                    st.session_state.input_ruc_cliente.strip() or "-"
                ),
                "Razón Social/Nombre": (
                    st.session_state.input_razon_social_cliente.strip() or "-"
                ),
                "Número Cliente": (
                    st.session_state.input_numero_cliente.strip() or "-"
                ),
                "Correo Cliente": (
                    st.session_state.input_correo_cliente.strip() or "-"
                ),
                "Registrado por": (
                    st.session_state.usuario_actual
                    if st.session_state.rol_usuario == "Supervisor"
                    else st.session_state.rol_usuario
                ),
            }
            try:
                insertar_registro(nuevo_registro)
            except (psycopg2.Error, ValueError) as error:
                st.error(f"No se pudo guardar el registro: {error}")
            else:
                refrescar_registros_sesion()
                limpiar_formulario()
                notificar_exito("¡Se registró correctamente!")
                st.rerun()
    else:
        st.warning("Por favor, ingresa al menos la placa o identificador.")

if btn_eliminar and st.session_state.edit_id:
    try:
        eliminado = eliminar_registro(st.session_state.edit_id)
    except (psycopg2.Error, ValueError) as error:
        st.error(f"No se pudo eliminar el registro: {error}")
    else:
        if eliminado:
            refrescar_registros_sesion()
            limpiar_formulario()
            notificar_exito(
                "El registro se eliminó correctamente.", "eliminacion"
            )
            st.rerun()
        else:
            st.error("No se encontró el registro a eliminar.")
