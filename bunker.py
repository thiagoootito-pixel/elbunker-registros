import tkinter as tk
import tkinter.font as tkfont
from tkinter import messagebox, ttk, filedialog
import csv
import os
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from tkcalendar import DateEntry
from PIL import Image, ImageTk, ImageDraw

class GestorRegistros:
    def __init__(self, archivo):
        self.archivo = archivo
        self.columnas = ["ID", "Tipo", "Placa", "Servicio", "Pago", "Monto", "Fecha", "Lavador 1", "Lavador 2", "Motivo Gasto", "Precio Gasto"]
        if not os.path.exists(self.archivo):
            self.guardar_todos([])

    def leer_registros(self):
        try:
            with open(self.archivo, "r", newline="", encoding="utf-8") as file:
                registros = list(csv.DictReader(file))
                for registro in registros:
                    gasto_antiguo = registro.pop("Gasto", "-")
                    if gasto_antiguo and gasto_antiguo != "-" and ": s/" in gasto_antiguo.lower():
                        separador = gasto_antiguo.lower().index(": s/")
                        registro.setdefault("Motivo Gasto", gasto_antiguo[:separador].strip())
                        registro.setdefault("Precio Gasto", f"S/ {gasto_antiguo[separador + 4:].strip()}")
                    registro.setdefault("Motivo Gasto", "-")
                    registro.setdefault("Precio Gasto", "-")
                return registros
        except (FileNotFoundError, KeyError): 
            return []

    def guardar_todos(self, registros):
        with open(self.archivo, "w", newline="", encoding="utf-8") as file:
            writer = csv.DictWriter(file, fieldnames=self.columnas)
            writer.writeheader()
            writer.writerows(registros)

    def agregar_registro(self, registro_dict):
        id_unico = datetime.now().strftime("%Y%m%d%H%M%S%f")
        registro_dict["ID"] = id_unico
        registros = self.leer_registros()
        registros.append(registro_dict)
        self.guardar_todos(registros)


class App:
    def __init__(self, root, gestor):
        self.gestor = gestor
        self.root = root
        self.root.title("Sistema Autolavado Pro v13.5 - The Bunker (Red & Black)")
        
        self.root.geometry("1450x780") 
        self.root.configure(bg="#1A1A1A") 
        self.password_maestra = "281114"
        
        # Archivos de configuración
        self.archivo_lavadores = "lavadores.txt"
        self.archivo_vehiculos = "vehiculos.txt"
        self.archivo_servicios = "servicios.txt"
        
        # Estados de la UI
        self.es_admin = False
        self.panel_admin_derecho = None  
        self.mostrar_segundo_lavador = False 
        self.mostrar_slots_gasto = False 

        # Cargar listas iniciales
        self.lista_lavadores = self.cargar_lista(
            self.archivo_lavadores,
            [
                "Ismael", "Eymar", "Antony", "Anderson", "Alvaro",
                "Jeferson", "Zara", "Dylan", "Wender", "Eduardo",
                "Javier", "Eglimar", "Gabriel", "Jose", "Jose elis",
                "Oscar", "Renzo", "Renzo V.", "Sergio", "Kelvin",
                "Merwin", "William", "Pedro", "Felix", "Marco",
            ],
        )
        self.lista_vehiculos = self.cargar_lista(self.archivo_vehiculos, ["Sedán", "SUV", "Camioneta", "Moto"])
        self.lista_servicios = self.cargar_lista(
            self.archivo_servicios,
            [
                "Lavado Básico",
                "Lavado Completo",
                "Encerado",
                "Motor",
                "L. Basico+Pulido",
                "L. Basico+Motor",
                "L. Basico+Chasis,Motor,Cera",
                "Tratamiento de cerámico",
            ],
        )

        # Caché de imágenes de botones
        self.imagenes_botones = {}

        self.estilo_input = {
            "highlightbackground": "#444444", 
            "highlightcolor": "#E74C3C", 
            "highlightthickness": 2, 
            "bd": 0, "relief": "flat",
            "bg": "#FFFFFF", "fg": "#000000"
        }

        self.setup_ui()
        self.root.bind("<Button-1>", self.clic_fuera)
        self.root.after(100, self.filtrar_por_fecha)

    def cargar_lista(self, archivo, base):
        if not os.path.exists(archivo):
            self.guardar_lista_archivo(archivo, base)
            return base
        with open(archivo, "r", encoding="utf-8") as f:
            return [line.strip() for line in f if line.strip()]

    def guardar_lista_archivo(self, archivo, lista):
        with open(archivo, "w", encoding="utf-8") as f:
            for item in lista:
                f.write(f"{item}\n")

    def crear_boton_redondeado(self, ancho, alto, radio, color_hex, texto, comando, master, fg_color="white", fuente=("Arial", 9, "bold")):
        key = f"{ancho}_{alto}_{radio}_{color_hex}_{texto.replace(' ', '_')}"
        
        img = Image.new("RGBA", (ancho, alto), (0, 0, 0, 0))
        draw = ImageDraw.Draw(img)
        draw.rounded_rectangle([0, 0, ancho - 1, alto - 1], radius=radio, fill=color_hex)
        
        tk_img = ImageTk.PhotoImage(img)
        self.imagenes_botones[key] = tk_img 
        
        btn = tk.Button(master, image=tk_img, text=texto, compound="center", fg=fg_color, font=fuente,
                        bd=0, bg=master["bg"], activebackground=master["bg"], cursor="hand2", command=comando)
        return btn

    def setup_ui(self):
        style = ttk.Style()
        style.theme_use("clam")

        self.tree_zoom_levels = (100, 110, 120, 130)
        self.tree_zoom_index = 0
        self.tree_base_rowheight = 25
        default_font = tkfont.nametofont("TkDefaultFont")
        self.tree_base_font_size = default_font.actual("size")
        self.tree_base_heading_size = 9
        self.tree_font = tkfont.Font(self.root, family=default_font.actual("family"), size=self.tree_base_font_size)
        self.tree_heading_font = tkfont.Font(self.root, family="Arial", size=self.tree_base_heading_size, weight="bold")
        self.tree_style = style
        
        # Visor de registros en Blanco nítido
        style.configure("Treeview", background="#FFFFFF", foreground="#000000", font=self.tree_font, rowheight=self.tree_base_rowheight, fieldbackground="#FFFFFF", bordercolor="#D3D3D3", borderwidth=1)
        style.layout("Treeview", [('Treeview.treearea', {'sticky': 'nswe'})]) 
        style.map("Treeview", background=[('selected', '#E74C3C')], foreground=[('selected', 'white')])
        style.configure("Treeview.Heading", background="#111111", foreground="white", borderwidth=1, font=self.tree_heading_font)

        style.configure("TNotebook", background="#111111", borderwidth=0)
        style.configure("TNotebook.Tab", background="#1A1A1A", foreground="#888888", padding=[10, 5], font=("Arial", 9, "bold"))
        style.map("TNotebook.Tab", background=[("selected", "#E74C3C")], foreground=[("selected", "white")])

        # --- Frame Superior ---
        self.frame_top = tk.Frame(self.root, bg="#1A1A1A", padx=15, pady=15)
        self.frame_top.pack(fill="x")

        self.form_frame = tk.Frame(self.frame_top, bg="#1A1A1A")
        self.form_frame.pack(side="left", fill="x", expand=True)

        tk.Label(self.form_frame, text="NUEVO SERVICIO", bg="#1A1A1A", fg="#E74C3C", font=("Arial", 10, "bold")).grid(row=0, column=0, columnspan=3, sticky="w", pady=(0,10))
        
        self.lbl_status_admin = tk.Label(self.form_frame, text="🔒 MODO OPERADOR", bg="#1A1A1A", fg="#888888", font=("Arial", 10, "bold"))
        self.lbl_status_admin.grid(row=0, column=4, columnspan=2, sticky="w", pady=(0,10))

        # Campos Estables
        campos = [("Tipo de Vehículo", 0), ("Placa", 1), ("Servicio", 2), ("Método Pago", 3), ("Monto Recibido", 4)]
        for texto, col in campos:
            tk.Label(self.form_frame, text=texto, bg="#1A1A1A", fg="#CCCCCC", font=("Arial", 9)).grid(row=1, column=col, sticky="w", padx=5)

        # Contenedor Lavador
        frame_lbl_lavador = tk.Frame(self.form_frame, bg="#1A1A1A")
        frame_lbl_lavador.grid(row=1, column=5, sticky="w", padx=5)
        tk.Label(frame_lbl_lavador, text="Lavador 1 ", bg="#1A1A1A", fg="#CCCCCC", font=("Arial", 9)).pack(side="left")
        self.btn_add_lavador = tk.Button(frame_lbl_lavador, text="[+]", command=self.toggle_slot_lavador, bg="#1A1A1A", fg="#2ECC71", font=("Arial", 9, "bold"), bd=0, cursor="hand2")
        self.btn_add_lavador.pack(side="left")

        # Contenedor Gasto
        frame_lbl_gasto = tk.Frame(self.form_frame, bg="#1A1A1A")
        frame_lbl_gasto.grid(row=1, column=6, sticky="w", padx=5)
        tk.Label(frame_lbl_gasto, text="Gasto ", bg="#1A1A1A", fg="#CCCCCC", font=("Arial", 9)).pack(side="left")
        self.btn_add_gasto = tk.Button(frame_lbl_gasto, text="[+]", command=self.toggle_slots_gasto, bg="#1A1A1A", fg="#E67E22", font=("Arial", 9, "bold"), bd=0, cursor="hand2")
        self.btn_add_gasto.pack(side="left")

        # Inputs Principales
        self.tipo_carro = ttk.Combobox(self.form_frame, values=self.lista_vehiculos, width=13, state="normal")
        self.placa = tk.Entry(self.form_frame, width=12, **self.estilo_input)
        self.servicio = ttk.Combobox(self.form_frame, values=self.lista_servicios, width=15, state="readonly")
        self.pago = ttk.Combobox(self.form_frame, values=["Efectivo", "Tarjeta", "Yape/Plin"], width=10, state="readonly")
        self.monto = tk.Entry(self.form_frame, width=8, **self.estilo_input)
        self.cb_lavador1 = ttk.Combobox(self.form_frame, values=self.lista_lavadores, width=14, state="readonly")
        
        self.tipo_carro.grid(row=2, column=0, padx=4, pady=5, sticky="n")
        self.placa.grid(row=2, column=1, padx=4, pady=5, sticky="n")
        self.servicio.grid(row=2, column=2, padx=4, pady=5, sticky="n")
        self.pago.grid(row=2, column=3, padx=4, pady=5, sticky="n")
        self.monto.grid(row=2, column=4, padx=4, pady=5, sticky="n")
        self.cb_lavador1.grid(row=2, column=5, padx=4, pady=5, sticky="n")

        # Slots dinámicos secundarios
        self.cb_lavador2 = ttk.Combobox(self.form_frame, values=self.lista_lavadores, width=14, state="readonly")
        
        self.frame_gasto_desplegable = tk.Frame(self.form_frame, bg="#1A1A1A")
        self.lbl_tag_motivo = tk.Label(self.frame_gasto_desplegable, text="Motivo Gasto: ", bg="#1A1A1A", fg="#E67E22", font=("Arial", 9, "bold"))
        self.lbl_tag_motivo.grid(row=0, column=0, sticky="e", pady=2)
        self.entry_motivo_gasto = tk.Entry(self.frame_gasto_desplegable, width=12, **self.estilo_input)
        self.entry_motivo_gasto.grid(row=0, column=1, sticky="w", pady=2)

        self.lbl_tag_costo = tk.Label(self.frame_gasto_desplegable, text="Precio Gasto: ", bg="#1A1A1A", fg="#E67E22", font=("Arial", 9, "bold"))
        self.lbl_tag_costo.grid(row=1, column=0, sticky="e", pady=2)
        self.entry_costo_gasto = tk.Entry(self.frame_gasto_desplegable, width=8, **self.estilo_input)
        self.entry_costo_gasto.grid(row=1, column=1, sticky="w", pady=2)

        # Contenedor Botones de Operación (Agregar, Editar, Eliminar)
        self.btn_frame = tk.Frame(self.form_frame, bg="#1A1A1A")
        self.btn_frame.grid(row=2, column=7, padx=15, rowspan=2, sticky="n")
        
        self.btn_agregar = self.crear_boton_redondeado(85, 28, 10, "#E74C3C", "✚ Agregar", self.ejecutar_alta, self.btn_frame)
        self.btn_agregar.pack(side="left", padx=3)
        
        self.btn_actualizar = self.crear_boton_redondeado(90, 28, 10, "#555555", "🔄 Actualizar", self.actualizar_registro, self.btn_frame)
        self.btn_actualizar.pack(side="left", padx=3)
        
        self.btn_eliminar = self.crear_boton_redondeado(85, 28, 10, "#555555", "🗑 Eliminar", self.eliminar_registro, self.btn_frame)
        self.btn_eliminar.pack(side="left", padx=3)

        # Panel Derecho Superior (Contenedor exclusivo del Logo)
        self.right_top_frame = tk.Frame(self.frame_top, bg="#1A1A1A")
        self.right_top_frame.pack(side="right", anchor="ne", padx=10)

        try:
            if os.path.exists("bunker_logo.png"):
                img = Image.open("bunker_logo.png")
                img = img.resize((210, 90), Image.Resampling.LANCZOS)
                self.logo_img = ImageTk.PhotoImage(img)
                logo_label = tk.Label(self.right_top_frame, image=self.logo_img, bg="#1A1A1A")
                logo_label.pack(anchor="ne")
        except Exception:
            pass

        # Frame Central (Filtros y Herramientas Globales)
        self.frame_mid = tk.Frame(self.root, bg="#111111", pady=10, padx=20)
        self.frame_mid.pack(fill="x", padx=10)

        tk.Label(self.frame_mid, text="🔍 Buscar:", bg="#111111", fg="white").pack(side="left", padx=5)
        self.buscar_entry = tk.Entry(self.frame_mid, **self.estilo_input)
        self.buscar_entry.pack(side="left", padx=5)
        self.buscar_entry.bind("<KeyRelease>", lambda e: self.buscar_registro())

        tk.Label(self.frame_mid, text="📅 Ver Fecha:", bg="#111111", fg="white").pack(side="left", padx=15)
        self.calendario = DateEntry(self.frame_mid, date_pattern="dd/mm/yyyy", state="readonly")
        self.calendario.pack(side="left", padx=5)
        self.calendario.bind("<<DateEntrySelected>>", lambda e: self.filtrar_por_fecha())

        self.frame_herramientas_der = tk.Frame(self.frame_mid, bg="#111111")
        self.frame_herramientas_der.pack(side="right")

        self.btn_excel = self.crear_boton_redondeado(130, 28, 10, "#16A085", "Generar Excel ✨", self.exportar_excel, self.frame_herramientas_der)
        self.btn_excel.pack(side="left", padx=5)

        self.btn_admin_mode = self.crear_boton_redondeado(160, 28, 10, "#E67E22", "🔑 Modo Administrador", self.toggle_modo_administrador, self.frame_herramientas_der)
        self.btn_admin_mode.pack(side="left", padx=5)

        # Contenedor Inferior Principal
        self.main_lower_container = tk.Frame(self.root, bg="#1A1A1A")
        self.main_lower_container.pack(pady=10, padx=10, fill="both", expand=True)

        self.table_container = tk.Frame(self.main_lower_container, bg="#FFFFFF")
        self.table_container.pack(side="left", fill="both", expand=True)

        columnas_visibles = ("Tipo", "Placa", "Servicio", "Pago", "Monto", "Fecha", "Lavador 1", "Lavador 2", "Motivo Gasto", "Precio Gasto")
        self.tree = ttk.Treeview(self.table_container, columns=("ID",) + columnas_visibles, show="headings")

        self.tree_zoom_frame = tk.Frame(self.table_container, bg="#FFFFFF")
        self.tree_zoom_frame.pack(side="bottom", fill="x")
        self.tree_zoom_controls = tk.Frame(self.tree_zoom_frame, bg="#FFFFFF")
        self.tree_zoom_controls.pack(side="right", padx=6, pady=3)
        tk.Button(self.tree_zoom_controls, text="[ - ]", command=lambda: self.ajustar_zoom_tree(-1), bg="#FFFFFF", fg="#222222", activebackground="#EEEEEE", bd=0, font=("Arial", 9, "bold"), cursor="hand2").pack(side="left", padx=3)
        self.lbl_tree_zoom = tk.Label(self.tree_zoom_controls, text="100%", bg="#FFFFFF", fg="#222222", font=("Arial", 9, "bold"), width=5)
        self.lbl_tree_zoom.pack(side="left", padx=3)
        tk.Button(self.tree_zoom_controls, text="[ + ]", command=lambda: self.ajustar_zoom_tree(1), bg="#FFFFFF", fg="#222222", activebackground="#EEEEEE", bd=0, font=("Arial", 9, "bold"), cursor="hand2").pack(side="left", padx=3)
        self.tree.bind("<Control-MouseWheel>", self.zoom_tree_rueda)
        
        scrollbar_v = ttk.Scrollbar(self.table_container, orient="vertical", command=self.tree.yview)
        scrollbar_h = ttk.Scrollbar(self.table_container, orient="horizontal", command=self.tree.xview)
        
        self.tree.configure(yscrollcommand=scrollbar_v.set, xscrollcommand=scrollbar_h.set)
        
        scrollbar_v.pack(side="right", fill="y")
        scrollbar_h.pack(side="bottom", fill="x")
        self.tree.pack(side="left", fill="both", expand=True)

        for col in ("ID",) + columnas_visibles:
            self.tree.heading(col, text=col)
            self.tree.column(col, width=130, anchor="center")
        
        self.tree["displaycolumns"] = columnas_visibles
        self.tree.bind("<<TreeviewSelect>>", self.seleccionar_registro)

    def ajustar_zoom_tree(self, cambio):
        self.tree_zoom_index = min(max(self.tree_zoom_index + cambio, 0), len(self.tree_zoom_levels) - 1)
        zoom = self.tree_zoom_levels[self.tree_zoom_index]
        factor = zoom / 100
        self.tree_font.configure(size=round(self.tree_base_font_size * factor))
        self.tree_heading_font.configure(size=round(self.tree_base_heading_size * factor))
        self.tree_style.configure("Treeview", rowheight=round(self.tree_base_rowheight * factor))
        self.lbl_tree_zoom.config(text=f"{zoom}%")

    def zoom_tree_rueda(self, event):
        self.ajustar_zoom_tree(1 if event.delta > 0 else -1)
        return "break"

    def toggle_slot_lavador(self):
        if not self.mostrar_segundo_lavador:
            self.cb_lavador2.grid(row=3, column=5, padx=4, pady=(0, 5), sticky="n")
            self.btn_add_lavador.config(text="[-]", fg="#E74C3C")
            self.mostrar_segundo_lavador = True
        else:
            self.cb_lavador2.grid_forget()
            self.cb_lavador2.set("")
            self.btn_add_lavador.config(text="[+]", fg="#2ECC71")
            self.mostrar_segundo_lavador = False

    def toggle_slots_gasto(self):
        if not self.mostrar_slots_gasto:
            self.frame_gasto_desplegable.grid(row=2, column=6, padx=4, pady=2, rowspan=2, sticky="nw")
            self.entry_motivo_gasto.focus_set()
            self.btn_add_gasto.config(text="[-]", fg="#E74C3C")
            self.mostrar_slots_gasto = True
        else:
            self.frame_gasto_desplegable.grid_forget()
            self.entry_motivo_gasto.delete(0, tk.END)
            self.entry_costo_gasto.delete(0, tk.END)
            self.btn_add_gasto.config(text="[+]", fg="#E67E22")
            self.mostrar_slots_gasto = False

    def mostrar_panel_admin(self):
        self.panel_admin_derecho = tk.Frame(self.main_lower_container, bg="#111111", width=360, padx=10, pady=5)
        self.panel_admin_derecho.pack(side="right", before=self.table_container, fill="both", padx=(10, 0))
        self.panel_admin_derecho.pack_propagate(False) 

        self.notebook = ttk.Notebook(self.panel_admin_derecho)
        self.notebook.pack(fill="both", expand=True)

        # TAB 1: CAJA
        self.tab_ganancias = tk.Frame(self.notebook, bg="#111111", padx=12, pady=12)
        self.notebook.add(self.tab_ganancias, text="📊 Caja")

        lbl_seccion_ingresos = tk.Label(self.tab_ganancias, text="💰 BALANCE DE INGRESOS", bg="#111111", fg="#2ECC71", font=("Arial", 10, "bold"))
        lbl_seccion_ingresos.pack(anchor="w", pady=(5, 2))
        
        frame_box_dia = tk.LabelFrame(self.tab_ganancias, text=" Del Día ", bg="#111111", fg="white", font=("Arial", 9, "bold"), padx=10, pady=10)
        frame_box_dia.pack(fill="x", pady=5)
        self.lbl_ganancia_dia_admin = tk.Label(frame_box_dia, text="S/ 0.00", bg="#111111", fg="#2ECC71", font=("Arial", 18, "bold"))
        self.lbl_ganancia_dia_admin.pack()

        frame_box_mes = tk.LabelFrame(self.tab_ganancias, text=" Del Mes Actual ", bg="#111111", fg="white", font=("Arial", 9, "bold"), padx=10, pady=10)
        frame_box_mes.pack(fill="x", pady=5)
        self.lbl_ganancia_mes_admin = tk.Label(frame_box_mes, text="S/ 0.00", bg="#111111", fg="#F1C40F", font=("Arial", 18, "bold"))
        self.lbl_ganancia_mes_admin.pack()

        lbl_separador = tk.Label(self.tab_ganancias, text="---------------------------------------------------", bg="#111111", fg="#444444")
        lbl_separador.pack(pady=10)

        lbl_seccion_gastos = tk.Label(self.tab_ganancias, text="📉 BALANCE DE GASTOS", bg="#111111", fg="#E74C3C", font=("Arial", 10, "bold"))
        lbl_seccion_gastos.pack(anchor="w", pady=(0, 2))

        frame_gasto_dia = tk.LabelFrame(self.tab_ganancias, text=" Del Día ", bg="#111111", fg="white", font=("Arial", 9, "bold"), padx=10, pady=10)
        frame_gasto_dia.pack(fill="x", pady=5)
        self.lbl_gasto_dia_admin = tk.Label(frame_gasto_dia, text="S/ 0.00", bg="#111111", fg="#E74C3C", font=("Arial", 18, "bold"))
        self.lbl_gasto_dia_admin.pack()

        frame_gasto_mes = tk.LabelFrame(self.tab_ganancias, text=" Del Mes Actual ", bg="#111111", fg="white", font=("Arial", 9, "bold"), padx=10, pady=10)
        frame_gasto_mes.pack(fill="x", pady=5)
        self.lbl_gasto_mes_admin = tk.Label(frame_gasto_mes, text="S/ 0.00", bg="#111111", fg="#C0392B", font=("Arial", 18, "bold"))
        self.lbl_gasto_mes_admin.pack()

        # TAB 2: PERSONAL
        self.tab_personal = tk.Frame(self.notebook, bg="#111111", padx=10, pady=15)
        self.notebook.add(self.tab_personal, text="👥 Personal")

        tk.Label(self.tab_personal, text="CONTROL DE LAVADORES", bg="#111111", fg="#E67E22", font=("Arial", 11, "bold")).pack(anchor="w", pady=(0, 10))
        
        frame_lista = tk.Frame(self.tab_personal, bg="#111111")
        frame_lista.pack(fill="both", expand=True, pady=5)

        self.listbox_personal = tk.Listbox(frame_lista, font=("Arial", 10), bd=0, highlightthickness=0, bg="#222222", fg="white", selectbackground="#E74C3C")
        self.listbox_personal.pack(side="left", fill="both", expand=True)
        
        scroll_p = ttk.Scrollbar(frame_lista, orient="vertical", command=self.listbox_personal.yview)
        self.listbox_personal.configure(yscrollcommand=scroll_p.set)
        scroll_p.pack(side="right", fill="y")
        
        frame_acciones_p = tk.Frame(self.tab_personal, bg="#111111", pady=10)
        frame_acciones_p.pack(fill="x")

        self.entry_nuevo_lavador = tk.Entry(frame_acciones_p, font=("Arial", 10), **self.estilo_input)
        self.entry_nuevo_lavador.pack(fill="x", pady=5)
        
        tk.Button(frame_acciones_p, text="✚ Registrar Lavador", command=self.admin_agregar_lavador, bg="#27AE60", fg="white", font=("Arial", 9, "bold"), bd=0, pady=4, cursor="hand2").pack(fill="x", pady=2)
        tk.Button(frame_acciones_p, text="🗑 Dar de Baja Seleccionado", command=self.admin_eliminar_lavador, bg="#C0392B", fg="white", font=("Arial", 9, "bold"), bd=0, pady=4, cursor="hand2").pack(fill="x", pady=2)

        # TAB 3: CONFIGURACIONES
        self.tab_config = tk.Frame(self.notebook, bg="#111111", padx=10, pady=10)
        self.notebook.add(self.tab_config, text="⚙️ Ajustes")

        lbl_v = tk.Label(self.tab_config, text="🚗 TIPOS DE VEHÍCULOS", bg="#111111", fg="#E74C3C", font=("Arial", 9, "bold"))
        lbl_v.pack(anchor="w", pady=(5,2))
        
        frame_v = tk.Frame(self.tab_config, bg="#111111")
        frame_v.pack(fill="both", expand=True)
        self.listbox_vehiculos = tk.Listbox(frame_v, font=("Arial", 9), bd=0, bg="#222222", fg="white", selectbackground="#E74C3C", height=5)
        self.listbox_vehiculos.pack(side="left", fill="both", expand=True)
        scroll_v = ttk.Scrollbar(frame_v, orient="vertical", command=self.listbox_vehiculos.yview)
        self.listbox_vehiculos.configure(yscrollcommand=scroll_v.set)
        scroll_v.pack(side="right", fill="y")

        self.entry_nuevo_vehiculo = tk.Entry(self.tab_config, font=("Arial", 9), **self.estilo_input)
        self.entry_nuevo_vehiculo.pack(fill="x", pady=2)
        
        frame_btn_v = tk.Frame(self.tab_config, bg="#111111")
        frame_btn_v.pack(fill="x", pady=(0, 10))
        tk.Button(frame_btn_v, text="✚ Añadir", command=self.admin_agregar_vehiculo, bg="#27AE60", fg="white", font=("Arial", 8, "bold"), bd=0, width=10, cursor="hand2").pack(side="left", padx=2)
        tk.Button(frame_btn_v, text="🗑 Quitar", command=self.admin_eliminar_vehiculo, bg="#C0392B", fg="white", font=("Arial", 8, "bold"), bd=0, width=10, cursor="hand2").pack(side="left", padx=2)

        lbl_s = tk.Label(self.tab_config, text="🛠 TIPOS DE SERVICIOS", bg="#111111", fg="#2ECC71", font=("Arial", 9, "bold"))
        lbl_s.pack(anchor="w", pady=(5,2))

        frame_s = tk.Frame(self.tab_config, bg="#111111")
        frame_s.pack(fill="both", expand=True)
        self.listbox_servicios = tk.Listbox(frame_s, font=("Arial", 9), bd=0, bg="#222222", fg="white", selectbackground="#E74C3C", height=5)
        self.listbox_servicios.pack(side="left", fill="both", expand=True)
        scroll_s = ttk.Scrollbar(frame_s, orient="vertical", command=self.listbox_servicios.yview)
        self.listbox_servicios.configure(yscrollcommand=scroll_s.set)
        scroll_s.pack(side="right", fill="y")

        self.entry_nuevo_servicio = tk.Entry(self.tab_config, font=("Arial", 9), **self.estilo_input)
        self.entry_nuevo_servicio.pack(fill="x", pady=2)

        frame_btn_s = tk.Frame(self.tab_config, bg="#111111")
        frame_btn_s.pack(fill="x")
        tk.Button(frame_btn_s, text="✚ Añadir", command=self.admin_agregar_servicio, bg="#27AE60", fg="white", font=("Arial", 8, "bold"), bd=0, width=10, cursor="hand2").pack(side="left", padx=2)
        tk.Button(frame_btn_s, text="🗑 Quitar", command=self.admin_eliminar_servicio, bg="#C0392B", fg="white", font=("Arial", 8, "bold"), bd=0, width=10, cursor="hand2").pack(side="left", padx=2)

        self.actualizar_listbox_admin()
        self.calcular_ganancias_avanzadas()

    def actualizar_listbox_admin(self):
        self.listbox_personal.delete(0, tk.END)
        for lav in self.lista_lavadores:
            self.listbox_personal.insert(tk.END, f"  👤 {lav}")
        
        self.listbox_vehiculos.delete(0, tk.END)
        for veh in self.lista_vehiculos:
            self.listbox_vehiculos.insert(tk.END, f"  🚗 {veh}")
        
        self.listbox_servicios.delete(0, tk.END)
        for ser in self.lista_servicios:
            self.listbox_servicios.insert(tk.END, f"  🛠 {ser}")

    def admin_agregar_lavador(self):
        nombre = self.entry_nuevo_lavador.get().strip().title()
        if not nombre or nombre in self.lista_lavadores: return
        self.lista_lavadores.append(nombre)
        self.guardar_lista_archivo(self.archivo_lavadores, self.lista_lavadores)
        self.actualizar_listbox_admin()
        self.cb_lavador1["values"] = self.lista_lavadores
        self.cb_lavador2["values"] = self.lista_lavadores
        self.entry_nuevo_lavador.delete(0, tk.END)

    def admin_eliminar_lavador(self):
        sel = self.listbox_personal.curselection()
        if not sel: return
        self.lista_lavadores.pop(sel[0])
        self.guardar_lista_archivo(self.archivo_lavadores, self.lista_lavadores)
        self.actualizar_listbox_admin()
        self.cb_lavador1["values"] = self.lista_lavadores
        self.cb_lavador2["values"] = self.lista_lavadores

    def admin_agregar_vehiculo(self):
        vehiculo = self.entry_nuevo_vehiculo.get().strip().title()
        # CORREGIDO: se cambió 'vehicle' por 'vehiculo'
        if not vehiculo or vehiculo in self.lista_vehiculos: return
        self.lista_vehiculos.append(vehiculo)
        self.guardar_lista_archivo(self.archivo_vehiculos, self.lista_vehiculos)
        self.actualizar_listbox_admin()
        self.tipo_carro["values"] = self.lista_vehiculos
        self.entry_nuevo_vehiculo.delete(0, tk.END)

    def admin_eliminar_vehiculo(self):
        sel = self.listbox_vehiculos.curselection()
        if not sel: return
        self.lista_vehiculos.pop(sel[0])
        self.guardar_lista_archivo(self.archivo_vehiculos, self.lista_vehiculos)
        self.actualizar_listbox_admin()
        self.tipo_carro["values"] = self.lista_vehiculos

    def admin_agregar_servicio(self):
        servicio = self.entry_nuevo_servicio.get().strip().title()
        if not servicio or servicio in self.lista_servicios: return
        self.lista_servicios.append(servicio)
        self.guardar_lista_archivo(self.archivo_servicios, self.lista_servicios)
        self.actualizar_listbox_admin()
        self.servicio["values"] = self.lista_servicios
        self.entry_nuevo_servicio.delete(0, tk.END)

    def admin_eliminar_servicio(self):
        sel = self.listbox_servicios.curselection()
        if not sel: return
        self.lista_servicios.pop(sel[0])
        self.guardar_lista_archivo(self.archivo_servicios, self.lista_servicios)
        self.actualizar_listbox_admin()
        self.servicio["values"] = self.lista_servicios

    def ocultar_panel_admin(self):
        if self.panel_admin_derecho:
            self.panel_admin_derecho.destroy()
            self.panel_admin_derecho = None

    def toggle_modo_administrador(self):
        if not self.es_admin:
            self.win_pwd = tk.Toplevel(self.root)
            self.win_pwd.title("Acceso")
            self.win_pwd.geometry("320x180")
            self.win_pwd.configure(bg="#1A1A1A")
            self.win_pwd.resizable(False, False)
            self.win_pwd.grab_set() 
            
            x = self.root.winfo_x() + (self.root.winfo_width() // 2) - 160
            y = self.root.winfo_y() + (self.root.winfo_height() // 2) - 90
            self.win_pwd.geometry(f"+{x}+{y}")

            tk.Label(self.win_pwd, text="🔑 CLAVE DE ADMINISTRADOR", bg="#1A1A1A", fg="white", font=("Arial", 10, "bold")).pack(pady=15)
            self.entry_pwd = tk.Entry(self.win_pwd, show="*", justify="center", font=("Arial", 14), **self.estilo_input)
            self.entry_pwd.pack(pady=5, padx=40, fill="x")
            self.entry_pwd.focus_set()
            self.entry_pwd.bind("<Return>", lambda e: self.validar_login_admin())

            tk.Button(self.win_pwd, text="INGRESAR", command=self.validar_login_admin, bg="#E74C3C", fg="white", font=("Arial", 10, "bold"), bd=0, height=2, cursor="hand2").pack(pady=15, padx=40, fill="x")
        else:
            self.es_admin = False
            self.btn_admin_mode.destroy()
            self.btn_admin_mode = self.crear_boton_redondeado(160, 28, 10, "#E67E22", "🔑 Modo Administrador", self.toggle_modo_administrador, self.frame_herramientas_der)
            self.btn_admin_mode.pack(side="left", padx=5)
            
            self.lbl_status_admin.config(text="🔒 MODO OPERADOR", fg="#888888")
            
            for widget in self.btn_frame.winfo_children(): widget.pack_forget()
            self.btn_agregar = self.crear_boton_redondeado(85, 28, 10, "#E74C3C", "✚ Agregar", self.ejecutar_alta, self.btn_frame)
            self.btn_actualizar = self.crear_boton_redondeado(90, 28, 10, "#555555", "🔄 Actualizar", self.actualizar_registro, self.btn_frame)
            self.btn_eliminar = self.crear_boton_redondeado(85, 28, 10, "#555555", "🗑 Eliminar", self.eliminar_registro, self.btn_frame)
            self.btn_agregar.pack(side="left", padx=3)
            self.btn_actualizar.pack(side="left", padx=3)
            self.btn_eliminar.pack(side="left", padx=3)
            
            self.ocultar_panel_admin()

    def validar_login_admin(self):
        if self.entry_pwd.get() == self.password_maestra:
            self.es_admin = True
            self.win_pwd.destroy()
            
            self.btn_admin_mode.destroy()
            self.btn_admin_mode = self.crear_boton_redondeado(160, 28, 10, "#C0392B", "🔓 Cerrar Modo Admin", self.toggle_modo_administrador, self.frame_herramientas_der)
            self.btn_admin_mode.pack(side="left", padx=5)
            
            self.lbl_status_admin.config(text="🔓 MODO ADMINISTRADOR ACTIVO", fg="#2ECC71")
            
            for widget in self.btn_frame.winfo_children(): widget.pack_forget()
            self.btn_agregar = self.crear_boton_redondeado(85, 28, 10, "#E74C3C", "✚ Agregar", self.ejecutar_alta, self.btn_frame)
            self.btn_actualizar = self.crear_boton_redondeado(90, 28, 10, "#E74C3C", "🔄 Actualizar", self.actualizar_registro, self.btn_frame)
            self.btn_eliminar = self.crear_boton_redondeado(85, 28, 10, "#C0392B", "🗑 Eliminar", self.eliminar_registro, self.btn_frame)
            self.btn_agregar.pack(side="left", padx=3)
            self.btn_actualizar.pack(side="left", padx=3)
            self.btn_eliminar.pack(side="left", padx=3)

            self.mostrar_panel_admin()
        else:
            messagebox.showerror("Error", "Clave incorrecta")

    def calcular_ganancias_avanzadas(self):
        if self.es_admin and self.panel_admin_derecho:
            fecha_sel = self.calendario.get()
            mes_act = datetime.now().strftime("/%m/%Y")
            
            t_dia_ingresos = 0.0
            t_mes_ingresos = 0.0
            t_dia_gastos = 0.0
            t_mes_gastos = 0.0
            
            for r in self.gestor.leer_registros():
                try:
                    m = float(r["Monto"].replace("S/", "").strip())
                    if r["Fecha"].startswith(fecha_sel): t_dia_ingresos += m
                    if mes_act in r["Fecha"].split()[0]: t_mes_ingresos += m
                except Exception: pass
                
                try:
                    val_gasto = float(r.get("Precio Gasto", "-").replace("S/", "").strip())
                    if r["Fecha"].startswith(fecha_sel): t_dia_gastos += val_gasto
                    if mes_act in r["Fecha"].split()[0]: t_mes_gastos += val_gasto
                except (ValueError, AttributeError): pass

            self.lbl_ganancia_dia_admin.config(text=f"S/ {t_dia_ingresos:.2f}")
            self.lbl_ganancia_mes_admin.config(text=f"S/ {t_mes_ingresos:.2f}")
            self.lbl_gasto_dia_admin.config(text=f"S/ {t_dia_gastos:.2f}")
            self.lbl_gasto_mes_admin.config(text=f"S/ {t_mes_gastos:.2f}")

    def filtrar_por_fecha(self, event=None):
        fecha = self.calendario.get()
        self.tree.delete(*self.tree.get_children())
        for r in self.gestor.leer_registros():
            if r["Fecha"].startswith(fecha):
                valores_ordenados = [r.get(col, "") for col in self.gestor.columnas]
                self.tree.insert("", tk.END, values=valores_ordenados)
        self.calcular_ganancias_avanzadas()

    def ejecutar_alta(self):
        if not all([self.tipo_carro.get().strip(), self.placa.get().strip(), self.servicio.get(), self.pago.get(), self.monto.get().strip(), self.cb_lavador1.get()]):
            messagebox.showerror("Error", "Los campos del servicio y el Lavador 1 son obligatorios.")
            return

        try:
            monto_val = float(self.monto.get())
            if monto_val < 0: raise ValueError()
        except ValueError:
            messagebox.showerror("Error", "Monto recibido inválido.")
            return

        motivo_gasto = "-"
        precio_gasto = "-"
        if self.mostrar_slots_gasto:
            motivo = self.entry_motivo_gasto.get().strip()
            costo_raw = self.entry_costo_gasto.get().strip()
            if not motivo or not costo_raw:
                messagebox.showerror("Error", "Complete los campos de Motivo y Costo del gasto.")
                return
            try:
                costo_val = float(costo_raw)
                if costo_val < 0: raise ValueError()
                motivo_gasto = motivo
                precio_gasto = f"S/ {costo_val:.2f}"
            except ValueError:
                messagebox.showerror("Error", "El costo del gasto debe ser un número positivo.")
                return

        lavador_1 = self.cb_lavador1.get()
        lavador_2 = "-"
        if self.mostrar_segundo_lavador and self.cb_lavador2.get():
            if self.cb_lavador1.get() == self.cb_lavador2.get():
                messagebox.showerror("Error", "No repitas el mismo lavador.")
                return
            lavador_2 = self.cb_lavador2.get()

        d = {
            "ID": "", 
            "Tipo": self.tipo_carro.get().strip().title(), 
            "Placa": self.placa.get().upper().strip(),
            "Servicio": self.servicio.get(), 
            "Pago": self.pago.get(),
            "Monto": f"S/ {monto_val:.2f}", 
            "Fecha": self.calendario.get() + " " + datetime.now().strftime("%H:%M"),
            "Lavador 1": lavador_1,  
            "Lavador 2": lavador_2,  
            "Motivo Gasto": motivo_gasto,
            "Precio Gasto": precio_gasto
        }
        self.gestor.agregar_registro(d)
        self.filtrar_por_fecha()
        self.limpiar_entradas()
        messagebox.showinfo("Éxito", "Registro guardado correctamente.")

    def eliminar_registro(self):
        if not self.es_admin: return
        sel = self.tree.selection()
        if not sel: return
        if messagebox.askyesno("Confirmar", "¿Eliminar definitivamente?"):
            id_b = self.tree.item(sel[0], "values")[0]
            nuevos = [r for r in self.gestor.leer_registros() if r["ID"] != id_b]
            self.gestor.guardar_todos(nuevos)
            self.filtrar_por_fecha()
            self.limpiar_entradas()

    def actualizar_registro(self):
        if not self.es_admin: return
        sel = self.tree.selection()
        if not sel: return
        
        try:
            monto_val = float(self.monto.get())
            if monto_val < 0: raise ValueError()
        except ValueError:
            messagebox.showerror("Error", "Monto recibido inválido.")
            return

        motivo_gasto = "-"
        precio_gasto = "-"
        if self.mostrar_slots_gasto:
            motivo = self.entry_motivo_gasto.get().strip()
            costo_raw = self.entry_costo_gasto.get().strip()
            if not motivo or not costo_raw:
                messagebox.showerror("Error", "Complete los campos de Motivo y Costo del gasto.")
                return
            try:
                costo_val = float(costo_raw)
                if costo_val < 0: raise ValueError()
                motivo_gasto = motivo
                precio_gasto = f"S/ {costo_val:.2f}"
            except ValueError:
                messagebox.showerror("Error", "El costo del gasto debe ser un número válido.")
                return

        lavador_1 = self.cb_lavador1.get()
        lavador_2 = "-"
        if self.mostrar_segundo_lavador and self.cb_lavador2.get():
            lavador_2 = self.cb_lavador2.get()

        id_a = self.tree.item(sel[0], "values")[0]
        regs = self.gestor.leer_registros()
        for r in regs:
            if r["ID"] == id_a:
                r.update({
                    "Tipo": self.tipo_carro.get().strip().title(), 
                    "Placa": self.placa.get().upper().strip(), 
                    "Servicio": self.servicio.get(), 
                    "Pago": self.pago.get(), 
                    "Monto": f"S/ {monto_val:.2f}", 
                    "Lavador 1": lavador_1, 
                    "Lavador 2": lavador_2, 
                    "Motivo Gasto": motivo_gasto,
                    "Precio Gasto": precio_gasto
                })
                break
        self.gestor.guardar_todos(regs)
        self.filtrar_por_fecha()
        self.limpiar_entradas()
        messagebox.showinfo("Éxito", "Registro actualizado correctamente.")

    def buscar_registro(self):
        q = self.buscar_entry.get().lower().strip()
        if not q: 
            self.filtrar_por_fecha()
            return
        self.tree.delete(*self.tree.get_children())
        for r in self.gestor.leer_registros():
            if any(q in str(v).lower() for v in r.values()):
                valores_ordenados = [r.get(col, "") for col in self.gestor.columnas]
                self.tree.insert("", tk.END, values=valores_ordenados)
        self.calcular_ganancias_avanzadas()

    def exportar_excel(self):
        registros = [self.tree.item(i, "values") for i in self.tree.get_children()]
        if not registros: 
            messagebox.showwarning("Vacío", "No hay datos para exportar")
            return
            
        fecha_hoy = datetime.now().strftime("%d-%m-%Y")
        archivo = filedialog.asksaveasfilename(
            defaultextension=".xlsx", 
            filetypes=[("Excel files", "*.xlsx")],
            initialfile=f"Reporte_{fecha_hoy}.xlsx"
        )
        
        if archivo:
            wb = Workbook()
            ws = wb.active
            ws.title = "Reporte de Lavado"
            
            encabezados = ["TIPO", "PLACA", "SERVICIO", "PAGO", "MONTO", "FECHA", "LAVADOR 1", "LAVADOR 2", "MOTIVO GASTO", "PRECIO GASTO"]
            ws.append(encabezados)
            
            fill_header = PatternFill(start_color="111111", end_color="111111", fill_type="solid")
            font_header = Font(color="FFFFFF", bold=True, size=12)
            alignment_center = Alignment(horizontal="center", vertical="center")
            
            # Definición de bordes más definidos (color oscuro para que se note bien)
            border_thin = Border(
                left=Side(style='thin', color='808080'),
                right=Side(style='thin', color='808080'),
                top=Side(style='thin', color='808080'),
                bottom=Side(style='thin', color='808080')
            )
            
            for cell in ws[1]:
                cell.fill = fill_header
                cell.font = font_header
                cell.alignment = alignment_center
                cell.border = border_thin
            
            for r in registros:
                ws.append(list(r)[1:])
                
            for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
                for cell in row:
                    cell.alignment = alignment_center
                    cell.border = border_thin  # Aplica bordes a todas las celdas de datos
                    if cell.row % 2 == 0:
                        cell.fill = PatternFill(start_color="F9F9F9", end_color="F9F9F9", fill_type="solid")

            for col in ws.columns:
                max_length = 0
                column = col[0].column_letter
                for cell in col:
                    try:
                        if len(str(cell.value)) > max_length:
                            max_length = len(str(cell.value))
                    except Exception: 
                        pass
                ws.column_dimensions[column].width = max_length + 5
                
            wb.save(archivo)
            messagebox.showinfo("Éxito", "Reporte Premium generado correctamente.")

    def seleccionar_registro(self, event):
        sel = self.tree.selection()
        if sel:
            v = self.tree.item(sel[0], "values")
            if len(v) < 11: return
            
            self.tipo_carro.set(v[1])
            self.placa.delete(0, tk.END); self.placa.insert(0, v[2])
            self.servicio.set(v[3])
            self.pago.set(v[4])
            self.monto.delete(0, tk.END); self.monto.insert(0, v[5].replace("S/", "").strip())
            
            motivo_gasto = v[9]
            precio_gasto = v[10]
            if motivo_gasto != "-" or precio_gasto != "-":
                if not self.mostrar_slots_gasto:
                    self.toggle_slots_gasto()
                self.entry_motivo_gasto.delete(0, tk.END)
                self.entry_motivo_gasto.insert(0, motivo_gasto if motivo_gasto != "-" else "")
                self.entry_costo_gasto.delete(0, tk.END)
                self.entry_costo_gasto.insert(0, precio_gasto.replace("S/", "").strip() if precio_gasto != "-" else "")
            else:
                if self.mostrar_slots_gasto:
                    self.toggle_slots_gasto()

            l1 = v[7]
            l2 = v[8]
            
            self.cb_lavador1.set(l1)
            if l2 != "-":
                if not self.mostrar_segundo_lavador: self.toggle_slot_lavador()
                self.cb_lavador2.set(l2)
            else:
                if self.mostrar_segundo_lavador: self.toggle_slot_lavador()

    def limpiar_entradas(self):
        for i in [self.placa, self.monto, self.entry_motivo_gasto, self.entry_costo_gasto]: 
            if i.winfo_exists(): i.delete(0, tk.END)
            
        for i in [self.tipo_carro, self.servicio, self.pago, self.cb_lavador1, self.cb_lavador2]: 
            i.set("")
            
        if self.mostrar_segundo_lavador: self.toggle_slot_lavador()
        if self.mostrar_slots_gasto: self.toggle_slots_gasto()
        if self.tree.selection(): self.tree.selection_remove(self.tree.selection())

    def clic_fuera(self, event):
        if event.widget in [self.root, self.tree] and self.tree.identify_row(event.y) == "":
            self.limpiar_entradas()


if __name__ == "__main__":
    root = tk.Tk()
    gestor = GestorRegistros("registros_lavado.csv")
    app = App(root, gestor)
    root.mainloop()